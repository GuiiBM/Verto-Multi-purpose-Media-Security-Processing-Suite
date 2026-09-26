"""Motor do Conversor Universal.

Cada conversão é uma rota "formato de entrada -> formato de saída" resolvida
por um motor real, sempre local (nada sai da máquina):

  - Pillow (+pillow-heif)    imagens raster
  - potrace (potracer)       vetorização de imagem -> SVG/DXF/EPS. É o mesmo
                             algoritmo usado pelos conversores online (Convertio,
                             online-convert): traça as bordas em curvas Bézier.
  - scikit-image             modo "linha central" (esqueleto do traço), para
                             gravação a laser/CNC/plotter, onde se quer uma linha
                             só e não o contorno dos dois lados do traço
  - ezdxf                    escrita/leitura de DXF e renderização DXF -> SVG/PDF/PNG
  - PyMuPDF                  PDF, SVG, XPS, EPUB/MOBI/FB2/CBZ e extração dos
                             vetores do PDF (PDF/SVG/AI/EPS -> DXF sem rasterizar)
  - Ghostscript              EPS/PS (entrada) e EPS/PS (saída)
  - LibreOffice headless     documentos, planilhas, apresentações e desenhos
                             (CDR, VSD, WMF/EMF...), reaproveitando o isolamento
                             de perfil do app Office
  - FFmpeg                   áudio e vídeo
  - 7-Zip + zipfile/tarfile  arquivos compactados
  - fontTools                fontes (TTF/OTF/WOFF/WOFF2, inclusive trocando o
                             tipo de contorno TrueType <-> CFF)
  - trimesh                  modelos 3D (e 3D -> DXF com 3DFACE)
"""

import csv
import io
import json
import math
import os
import re
import secrets
import shutil
import subprocess
import tarfile
import tempfile
import threading
import zipfile
from pathlib import Path


class ConversionError(Exception):
    pass


# ---------------------------------------------------------------------------
# Catálogo de formatos
# ---------------------------------------------------------------------------

GROUPS = [
    ('image', 'Imagem', '🖼️'),
    ('vector', 'Vetor', '✒️'),
    ('cad', 'CAD', '📐'),
    ('document', 'Documento', '📄'),
    ('spreadsheet', 'Planilha', '📊'),
    ('presentation', 'Apresentação', '📽️'),
    ('ebook', 'eBook', '📚'),
    ('data', 'Dados', '🧾'),
    ('audio', 'Áudio', '🎵'),
    ('video', 'Vídeo', '🎬'),
    ('archive', 'Compactado', '📦'),
    ('font', 'Fonte', '🔤'),
    ('model3d', '3D', '🧊'),
]

FORMATS = {
    # Imagem (raster)
    'png': ('image', 'Portable Network Graphics'),
    'jpg': ('image', 'JPEG'),
    'jpeg': ('image', 'JPEG'),
    'jfif': ('image', 'JPEG (JFIF)'),
    'webp': ('image', 'WebP'),
    'bmp': ('image', 'Bitmap do Windows'),
    'dib': ('image', 'Bitmap do Windows (DIB)'),
    'gif': ('image', 'GIF (inclusive animado)'),
    'tiff': ('image', 'TIFF'),
    'tif': ('image', 'TIFF'),
    'ico': ('image', 'Ícone do Windows'),
    'icns': ('image', 'Ícone do macOS'),
    'heic': ('image', 'HEIC (fotos do iPhone)'),
    'heif': ('image', 'HEIF'),
    'avif': ('image', 'AVIF'),
    'tga': ('image', 'Targa'),
    'ppm': ('image', 'Portable Pixmap'),
    'pgm': ('image', 'Portable Graymap'),
    'pbm': ('image', 'Portable Bitmap'),
    'pnm': ('image', 'Portable Anymap'),
    'pcx': ('image', 'PC Paintbrush'),
    'psd': ('image', 'Photoshop (camada achatada)'),
    'dds': ('image', 'DirectDraw Surface'),
    'jp2': ('image', 'JPEG 2000'),
    'j2k': ('image', 'JPEG 2000 (codestream)'),
    'qoi': ('image', 'Quite OK Image'),
    'sgi': ('image', 'Silicon Graphics'),
    'xbm': ('image', 'X Bitmap'),
    # Vetor
    'svg': ('vector', 'Scalable Vector Graphics'),
    'eps': ('vector', 'Encapsulated PostScript'),
    'ps': ('vector', 'PostScript'),
    'ai': ('vector', 'Adobe Illustrator (PDF compatível)'),
    'cdr': ('vector', 'CorelDRAW'),
    'wmf': ('vector', 'Windows Metafile'),
    'emf': ('vector', 'Enhanced Metafile'),
    'odg': ('vector', 'Desenho OpenDocument'),
    'vsd': ('vector', 'Visio'),
    'vsdx': ('vector', 'Visio (XML)'),
    'wpg': ('vector', 'WordPerfect Graphics'),
    # CAD
    'dxf': ('cad', 'AutoCAD Drawing Exchange Format'),
    'dwg': ('cad', 'AutoCAD Drawing'),
    # Documento
    'pdf': ('document', 'Portable Document Format'),
    'docx': ('document', 'Word'),
    'doc': ('document', 'Word 97-2003'),
    'odt': ('document', 'Texto OpenDocument'),
    'rtf': ('document', 'Rich Text Format'),
    'txt': ('document', 'Texto puro'),
    'html': ('document', 'Página HTML'),
    'htm': ('document', 'Página HTML'),
    'md': ('document', 'Markdown'),
    'wpd': ('document', 'WordPerfect'),
    'wps': ('document', 'Microsoft Works'),
    'pages': ('document', 'Apple Pages'),
    # Planilha
    'xlsx': ('spreadsheet', 'Excel'),
    'xls': ('spreadsheet', 'Excel 97-2003'),
    'ods': ('spreadsheet', 'Planilha OpenDocument'),
    'csv': ('spreadsheet', 'Valores separados por vírgula'),
    'numbers': ('spreadsheet', 'Apple Numbers'),
    # Apresentação
    'pptx': ('presentation', 'PowerPoint'),
    'ppt': ('presentation', 'PowerPoint 97-2003'),
    'pps': ('presentation', 'Apresentação de slides PowerPoint'),
    'ppsx': ('presentation', 'Apresentação de slides PowerPoint'),
    'odp': ('presentation', 'Apresentação OpenDocument'),
    'key': ('presentation', 'Apple Keynote'),
    # eBook
    'epub': ('ebook', 'EPUB'),
    'mobi': ('ebook', 'Kindle MOBI'),
    'fb2': ('ebook', 'FictionBook'),
    'cbz': ('ebook', 'Quadrinhos (CBZ)'),
    'xps': ('ebook', 'XML Paper Specification'),
    'oxps': ('ebook', 'OpenXPS'),
    # Dados
    'json': ('data', 'JSON'),
    'yaml': ('data', 'YAML'),
    'yml': ('data', 'YAML'),
    'xml': ('data', 'XML'),
    'tsv': ('data', 'Valores separados por tabulação'),
    # Áudio
    'mp3': ('audio', 'MP3'),
    'wav': ('audio', 'WAV'),
    'flac': ('audio', 'FLAC (sem perdas)'),
    'aac': ('audio', 'AAC'),
    'm4a': ('audio', 'MPEG-4 Áudio'),
    'ogg': ('audio', 'Ogg Vorbis'),
    'oga': ('audio', 'Ogg Áudio'),
    'opus': ('audio', 'Opus'),
    'wma': ('audio', 'Windows Media Audio'),
    'aiff': ('audio', 'AIFF'),
    'aif': ('audio', 'AIFF'),
    'ac3': ('audio', 'Dolby Digital AC-3'),
    'amr': ('audio', 'AMR (gravação de voz)'),
    'mka': ('audio', 'Matroska Áudio'),
    'caf': ('audio', 'Core Audio'),
    'mp2': ('audio', 'MPEG Layer II'),
    # Vídeo
    'mp4': ('video', 'MPEG-4'),
    'mkv': ('video', 'Matroska'),
    'avi': ('video', 'AVI'),
    'mov': ('video', 'QuickTime'),
    'webm': ('video', 'WebM'),
    'flv': ('video', 'Flash Video'),
    'wmv': ('video', 'Windows Media Video'),
    'm4v': ('video', 'MPEG-4 Vídeo'),
    'mpg': ('video', 'MPEG'),
    'mpeg': ('video', 'MPEG'),
    '3gp': ('video', '3GPP (celular)'),
    'ts': ('video', 'MPEG Transport Stream'),
    'mts': ('video', 'AVCHD'),
    'm2ts': ('video', 'Blu-ray BDAV'),
    'ogv': ('video', 'Ogg Vídeo'),
    'vob': ('video', 'DVD Vídeo'),
    'asf': ('video', 'Advanced Systems Format'),
    # Compactado
    'zip': ('archive', 'ZIP'),
    '7z': ('archive', '7-Zip'),
    'rar': ('archive', 'RAR (somente leitura)'),
    'tar': ('archive', 'TAR'),
    'tar.gz': ('archive', 'TAR + Gzip'),
    'tgz': ('archive', 'TAR + Gzip'),
    'tar.bz2': ('archive', 'TAR + Bzip2'),
    'tbz2': ('archive', 'TAR + Bzip2'),
    'tar.xz': ('archive', 'TAR + XZ'),
    'txz': ('archive', 'TAR + XZ'),
    'gz': ('archive', 'Gzip'),
    'bz2': ('archive', 'Bzip2'),
    'xz': ('archive', 'XZ'),
    'iso': ('archive', 'Imagem de disco ISO'),
    'cab': ('archive', 'Gabinete do Windows'),
    'arj': ('archive', 'ARJ'),
    'lzh': ('archive', 'LHA/LZH'),
    'jar': ('archive', 'Java Archive'),
    'apk': ('archive', 'Pacote Android'),
    'deb': ('archive', 'Pacote Debian'),
    'rpm': ('archive', 'Pacote RPM'),
    # Fonte
    'ttf': ('font', 'TrueType'),
    'otf': ('font', 'OpenType (CFF)'),
    'woff': ('font', 'Web Open Font'),
    'woff2': ('font', 'Web Open Font 2'),
    # 3D
    'stl': ('model3d', 'Estereolitografia (impressão 3D)'),
    'obj': ('model3d', 'Wavefront OBJ'),
    'ply': ('model3d', 'Polygon File Format'),
    'off': ('model3d', 'Object File Format'),
    'glb': ('model3d', 'glTF binário'),
    '3mf': ('model3d', '3D Manufacturing Format'),
}

# Extensões que são apenas apelidos de outra (entram como entrada, nunca como saída)
ALIASES = {'jpeg': 'jpg', 'jfif': 'jpg', 'tif': 'tiff', 'htm': 'html', 'yml': 'yaml',
           'heif': 'heic', 'aif': 'aiff', 'oga': 'ogg', 'mpeg': 'mpg', 'dib': 'bmp',
           'tgz': 'tar.gz', 'tbz2': 'tar.bz2', 'txz': 'tar.xz', 'j2k': 'jp2',
           'pnm': 'ppm', 'pgm': 'ppm', 'pbm': 'ppm', 'pps': 'ppt', 'ppsx': 'pptx',
           'oxps': 'xps'}

COMPOUND_EXTS = ('tar.gz', 'tar.bz2', 'tar.xz')

RASTER_IN = [e for e, (g, _) in FORMATS.items() if g == 'image']
RASTER_OUT = ['png', 'jpg', 'webp', 'bmp', 'gif', 'tiff', 'ico', 'icns', 'heic', 'avif',
              'tga', 'ppm', 'pcx', 'jp2', 'qoi']
TRACE_OUT = ['dxf', 'svg', 'eps']
PAGE_RASTER_OUT = ['png', 'jpg', 'webp', 'tiff', 'bmp', 'gif']

PDFLIKE_IN = ['pdf', 'ai', 'eps', 'ps', 'svg', 'xps', 'oxps']
EBOOK_IN = ['epub', 'mobi', 'fb2', 'cbz']
DRAW_IN = ['cdr', 'wmf', 'emf', 'odg', 'vsd', 'vsdx', 'wpg']
DOC_IN = ['doc', 'docx', 'odt', 'rtf', 'txt', 'html', 'htm', 'md', 'wpd', 'wps', 'pages']
DOC_OUT = ['pdf', 'docx', 'doc', 'odt', 'rtf', 'txt', 'html', 'md', 'epub', 'png', 'jpg']
SHEET_IN = ['xlsx', 'xls', 'ods', 'csv', 'numbers']
TABLE_OUT = ['xlsx', 'xls', 'ods', 'csv', 'tsv', 'json', 'yaml', 'xml', 'html', 'pdf']
DATA_IN = ['json', 'yaml', 'yml', 'xml', 'tsv']
SLIDES_IN = ['pptx', 'ppt', 'pps', 'ppsx', 'odp', 'key']
SLIDES_OUT = ['pptx', 'ppt', 'odp', 'pdf', 'png', 'jpg']
AUDIO_IN = [e for e, (g, _) in FORMATS.items() if g == 'audio']
AUDIO_OUT = ['mp3', 'wav', 'flac', 'aac', 'm4a', 'ogg', 'opus', 'wma', 'aiff', 'ac3', 'amr', 'mp2']
VIDEO_IN = [e for e, (g, _) in FORMATS.items() if g == 'video']
VIDEO_OUT = ['mp4', 'mkv', 'avi', 'mov', 'webm', 'flv', 'wmv', 'm4v', 'mpg', '3gp', 'ts', 'ogv']
ARCHIVE_IN = [e for e, (g, _) in FORMATS.items() if g == 'archive']
ARCHIVE_OUT = ['zip', '7z', 'tar', 'tar.gz', 'tar.bz2', 'tar.xz']
FONT_EXTS = ['ttf', 'otf', 'woff', 'woff2']
MODEL_IN = ['stl', 'obj', 'ply', 'off', 'glb', '3mf']
MODEL_OUT = ['stl', 'obj', 'ply', 'off', 'glb', 'dxf']

# Limites de segurança para a vetorização (o potrace é Python puro: imagens
# gigantes demorariam minutos sem ganho visível de qualidade)
TRACE_MAX_SIDE = 2400
TRACE_MAX_SIDE_COLOR = 1400
TRACE_MIN_SIDE = 900


# ---------------------------------------------------------------------------
# Ferramentas externas opcionais
# ---------------------------------------------------------------------------

def _which(*names):
    for name in names:
        found = shutil.which(name)
        if found:
            return found
    return None


def _oda_converter():
    return _which('ODAFileConverter', 'ODAFileConverter.exe')


def _libredwg_dwg2dxf():
    return _which('dwg2dxf')


def tool_status():
    """Quais motores externos existem nesta máquina (usado para esconder rotas
    que não funcionariam e explicar o que instalar)."""
    return {
        'ffmpeg': bool(_which('ffmpeg')),
        'soffice': bool(_which('soffice', 'libreoffice')),
        'ghostscript': bool(_which('gs', 'gswin64c', 'gswin32c')),
        '7z': bool(_which('7z', '7za', '7zz')),
        'oda': bool(_oda_converter()),
        'dwg2dxf': bool(_libredwg_dwg2dxf()),
    }


# ---------------------------------------------------------------------------
# Tabela de rotas
# ---------------------------------------------------------------------------

def _build_routes():
    """Monta {entrada: {saída: [grupos de opções]}}. Os grupos de opções dizem à
    interface quais controles mostrar para aquela combinação."""
    tools = tool_status()
    routes = {}

    def add(ins, outs, opts=()):
        for i in ins:
            for o in outs:
                if o == i or ALIASES.get(i) == o:
                    continue
                routes.setdefault(i, {})[o] = list(opts)

    # Imagens
    for i in RASTER_IN:
        add([i], RASTER_OUT, ['raster'])
        add([i], ['pdf'], ['raster'])
        add([i], ['svg', 'eps'], ['trace'])
        add([i], ['dxf'], ['trace', 'dxf'])
    if tools['ffmpeg']:
        add(['gif'], ['mp4', 'webm', 'avi', 'mov'], ['video'])

    # PDF e vetores "tipo PDF" (abertos pelo PyMuPDF ou via Ghostscript)
    pdflike = [e for e in PDFLIKE_IN if e not in ('eps', 'ps') or tools['ghostscript']]
    for i in pdflike:
        add([i], PAGE_RASTER_OUT, ['pages'])
        add([i], ['dxf'], ['dxf'])
        add([i], ['svg', 'pdf'])
        if tools['ghostscript']:
            add([i], ['eps', 'ps'])
    add(['pdf', 'xps', 'oxps'], ['txt', 'html'])
    if tools['soffice']:
        add(['pdf'], ['docx', 'odt'])

    # eBooks
    for i in EBOOK_IN:
        add([i], ['pdf', 'txt', 'html'] + ['png', 'jpg'], ['pages'])
        routes[i]['pdf'] = []
        routes[i]['txt'] = []
        routes[i]['html'] = []
        if tools['soffice']:
            add([i], ['docx', 'odt', 'epub'])

    # CAD
    add(['dxf'], ['svg', 'pdf'], ['cadview'])
    add(['dxf'], ['png', 'jpg', 'webp', 'tiff'], ['cadview', 'pages'])
    routes['dxf']['dxf'] = ['dxfversion']
    if tools['ghostscript']:
        add(['dxf'], ['eps'], ['cadview'])
    if tools['oda'] or tools['dwg2dxf']:
        add(['dwg'], ['dxf'], ['dxfversion'])
        add(['dwg'], ['svg', 'pdf'], ['cadview'])
        add(['dwg'], ['png', 'jpg', 'webp', 'tiff'], ['cadview', 'pages'])
    if tools['oda']:
        add(['dxf'], ['dwg'], ['dwgversion'])

    if tools['soffice']:
        # Desenhos que só o LibreOffice Draw abre (CorelDRAW, Visio, WMF/EMF...)
        add(DRAW_IN, ['pdf', 'svg', 'dxf'], [])
        for i in DRAW_IN:
            routes[i]['dxf'] = ['dxf']
        add(DRAW_IN, PAGE_RASTER_OUT, ['pages'])
        if tools['ghostscript']:
            add(DRAW_IN, ['eps'])

        # Documentos, planilhas e apresentações
        add(DOC_IN, DOC_OUT)
        for i in DOC_IN:
            routes[i]['png'] = ['pages']
            routes[i]['jpg'] = ['pages']
        add(SHEET_IN + DATA_IN, TABLE_OUT)
        add(SLIDES_IN, SLIDES_OUT)
        for i in SLIDES_IN:
            routes[i]['png'] = ['pages']
            routes[i]['jpg'] = ['pages']
    else:
        # Sem LibreOffice ainda dá para converter dados e Markdown/HTML
        add(['csv', 'xlsx'] + DATA_IN, ['xlsx', 'csv', 'tsv', 'json', 'yaml', 'xml', 'html'])
        add(['md'], ['html', 'txt'])
        add(['html', 'htm'], ['md', 'txt'])

    # Áudio e vídeo
    if tools['ffmpeg']:
        add(AUDIO_IN, AUDIO_OUT, ['audio'])
        add(AUDIO_IN, ['mp4'], ['audio'])
        add(VIDEO_IN, VIDEO_OUT, ['video'])
        add(VIDEO_IN, AUDIO_OUT, ['audio'])
        add(VIDEO_IN, ['gif'], ['gif'])

    # Compactados
    archive_in = ARCHIVE_IN if tools['7z'] else ['zip', 'tar', 'tar.gz', 'tgz', 'tar.bz2',
                                                   'tbz2', 'tar.xz', 'txz', 'gz', 'bz2', 'xz', 'jar', 'apk']
    archive_out = ARCHIVE_OUT if tools['7z'] else [o for o in ARCHIVE_OUT if o != '7z']
    add(archive_in, archive_out)

    # Fontes e 3D
    add(FONT_EXTS, FONT_EXTS)
    add(MODEL_IN, MODEL_OUT)
    for i in MODEL_IN:
        if 'dxf' in routes.get(i, {}):
            routes[i]['dxf'] = ['dxfversion']

    return routes


def unavailable_notes():
    """Formatos conhecidos mas que dependem de uma ferramenta não instalada."""
    tools = tool_status()
    notes = {}
    if not (tools['oda'] or tools['dwg2dxf']):
        notes['dwg'] = ("DWG é um formato fechado da Autodesk. Para abrir/gerar DWG instale o "
                        "ODA File Converter (gratuito, opendesign.com/guestfiles/oda_file_converter) "
                        "ou o LibreDWG (dwg2dxf) - o Conversor detecta sozinho. Enquanto isso, "
                        "use DXF: todo programa CAD (AutoCAD, LibreCAD, QCAD, SolidWorks, "
                        "Fusion) abre DXF.")
    elif not tools['oda']:
        notes['dwg'] = ("Com o LibreDWG dá para ler DWG; para gerar DWG instale o ODA File "
                        "Converter (gratuito).")
    if not tools['soffice']:
        notes['docx'] = "Documentos precisam do LibreOffice instalado (o INSTALAR.py cuida disso)."
    if not tools['ffmpeg']:
        notes['mp4'] = "Áudio e vídeo precisam do FFmpeg instalado (o INSTALAR.py cuida disso)."
    if not tools['ghostscript']:
        notes['eps'] = "EPS/PS precisam do Ghostscript (sudo apt install ghostscript)."
    return notes


def catalog():
    routes = _build_routes()
    return {
        'groups': [{'id': g, 'label': label, 'icon': icon} for g, label, icon in GROUPS],
        'formats': {ext: {'group': g, 'label': label} for ext, (g, label) in FORMATS.items()},
        'aliases': ALIASES,
        'routes': routes,
        'unavailable': unavailable_notes(),
    }


# ---------------------------------------------------------------------------
# Utilidades gerais
# ---------------------------------------------------------------------------

def split_ext(filename):
    name = os.path.basename(filename or '')
    lower = name.lower()
    for comp in COMPOUND_EXTS:
        if lower.endswith('.' + comp):
            return name[:-(len(comp) + 1)], comp
    base, ext = os.path.splitext(name)
    return base, ext.lower().lstrip('.')


def _safe_base(name):
    name = re.sub(r'[\\/:*?"<>|\x00-\x1f]+', '_', name or '').strip(' .')
    return name[:120] or 'arquivo'


def _downloads_dir():
    d = str(Path.home() / "Downloads")
    os.makedirs(d, exist_ok=True)
    return d


def _unique_path(directory, filename):
    base, ext = split_ext(filename)
    ext = '.' + ext if ext else ''
    path = os.path.join(directory, filename)
    n = 1
    while os.path.exists(path):
        path = os.path.join(directory, f"{base}_{n}{ext}")
        n += 1
    return path


def _run(cmd, timeout=600, what='a ferramenta'):
    try:
        result = subprocess.run(cmd, capture_output=True, timeout=timeout)
    except subprocess.TimeoutExpired as e:
        raise ConversionError(f"{what} demorou demais (tempo limite excedido).") from e
    except FileNotFoundError as e:
        raise ConversionError(f"{what} não está instalado neste sistema.") from e
    return result


def _last_error_line(result, fallback):
    text = (result.stderr or b'').decode('utf-8', 'ignore').strip() or \
           (result.stdout or b'').decode('utf-8', 'ignore').strip()
    lines = [l for l in text.splitlines() if l.strip()]
    return lines[-1] if lines else fallback


def _opt_int(opts, key, default, lo=None, hi=None):
    try:
        v = int(float(opts.get(key, default)))
    except (TypeError, ValueError):
        v = default
    if lo is not None:
        v = max(lo, v)
    if hi is not None:
        v = min(hi, v)
    return v


def _opt_float(opts, key, default, lo=None, hi=None):
    raw = opts.get(key, default)
    try:
        v = float(str(raw).replace(',', '.'))
    except (TypeError, ValueError):
        v = default
    if v != v:  # NaN
        v = default
    if lo is not None:
        v = max(lo, v)
    if hi is not None:
        v = min(hi, v)
    return v


def _opt_bool(opts, key, default=False):
    raw = opts.get(key)
    if raw is None or raw == '':
        return default
    return str(raw).lower() in ('1', 'true', 'on', 'yes', 'sim')


def _hex(rgb):
    return '#%02x%02x%02x' % tuple(int(c) for c in rgb[:3])


# ---------------------------------------------------------------------------
# Detecção do formato real
# ---------------------------------------------------------------------------

def sniff_ext(path, claimed_ext):
    """Confia na extensão quando ela bate com o conteúdo; senão, descobre pelo
    conteúdo (ex.: .enc do WhatsApp que é JPEG, .dat que é PDF, arquivo sem extensão)."""
    with open(path, 'rb') as f:
        head = f.read(4096)

    def is_image():
        try:
            from PIL import Image
            _register_heif()
            with Image.open(path) as im:
                return (im.format or '').lower()
        except Exception:
            return None

    if claimed_ext in FORMATS:
        group = FORMATS[claimed_ext][0]
        if group != 'image' or claimed_ext in ('psd', 'dds') or is_image():
            return claimed_ext

    if head.startswith(b'%PDF'):
        return 'pdf'
    if head.startswith(b'%!PS'):
        return 'eps' if b'EPSF' in head[:64] else 'ps'
    img = is_image()
    if img:
        return {'jpeg': 'jpg', 'mpo': 'jpg', 'tiff': 'tiff', 'heif': 'heic'}.get(img, img if img in FORMATS else 'png')
    stripped = head.lstrip()
    if stripped.startswith(b'<svg') or (stripped.startswith(b'<?xml') and b'<svg' in head):
        return 'svg'
    if b'SECTION' in head[:512] and (b'HEADER' in head[:512] or b'ENTITIES' in head[:512]):
        return 'dxf'
    if re.match(rb'AC10\d\d', head[:6]):
        return 'dwg'
    if head.startswith(b'PK\x03\x04'):
        try:
            with zipfile.ZipFile(path) as zf:
                names = zf.namelist()
            if '[Content_Types].xml' in names:
                if any(n.startswith('word/') for n in names):
                    return 'docx'
                if any(n.startswith('xl/') for n in names):
                    return 'xlsx'
                if any(n.startswith('ppt/') for n in names):
                    return 'pptx'
                if any(n.startswith('3D/') for n in names):
                    return '3mf'
            if 'mimetype' in names:
                with zipfile.ZipFile(path) as zf:
                    mime = zf.read('mimetype').decode('ascii', 'ignore')
                if 'epub' in mime:
                    return 'epub'
                for key, ext in (('text', 'odt'), ('spreadsheet', 'ods'),
                                 ('presentation', 'odp'), ('graphics', 'odg')):
                    if key in mime:
                        return ext
        except zipfile.BadZipFile:
            pass
        return 'zip'
    if head.startswith(b'7z\xbc\xaf\x27\x1c'):
        return '7z'
    if head.startswith(b'Rar!'):
        return 'rar'
    if head.startswith(b'\x1f\x8b'):
        return 'gz'
    if head.startswith(b'BZh'):
        return 'bz2'
    if head.startswith(b'\xfd7zXZ'):
        return 'xz'
    if head.startswith(b'ID3') or head[:2] in (b'\xff\xfb', b'\xff\xf3'):
        return 'mp3'
    if head[:4] == b'RIFF' and head[8:12] == b'WAVE':
        return 'wav'
    if head[:4] == b'RIFF' and head[8:12] == b'AVI ':
        return 'avi'
    if head.startswith(b'fLaC'):
        return 'flac'
    if head.startswith(b'OggS'):
        return 'ogg'
    if head[4:8] == b'ftyp':
        brand = head[8:12]
        if brand in (b'M4A ', b'M4B '):
            return 'm4a'
        if brand in (b'heic', b'heix', b'mif1'):
            return 'heic'
        if brand == b'qt  ':
            return 'mov'
        return 'mp4'
    if head.startswith(b'\x1a\x45\xdf\xa3'):
        return 'webm' if b'webm' in head[:64] else 'mkv'
    if head.startswith(b'\xd0\xcf\x11\xe0'):
        try:
            from Office import office_engine
            sub = office_engine.sniff_ole_subtype(head + open(path, 'rb').read(65536))
            return {'document': 'doc', 'spreadsheet': 'xls', 'presentation': 'ppt'}.get(sub, 'doc')
        except Exception:
            return 'doc'
    if head.startswith(b'wOFF'):
        return 'woff'
    if head.startswith(b'wOF2'):
        return 'woff2'
    if head.startswith(b'OTTO'):
        return 'otf'
    if head.startswith(b'\x00\x01\x00\x00'):
        return 'ttf'
    if head.startswith(b'solid'):
        return 'stl'
    if stripped[:1] in (b'{', b'['):
        return 'json'
    if stripped.startswith(b'<?xml') or stripped.startswith(b'<'):
        return 'html' if b'<html' in head.lower() else 'xml'
    try:
        head.decode('utf-8')
        return 'txt'
    except UnicodeDecodeError:
        pass
    return claimed_ext or ''


# ---------------------------------------------------------------------------
# Imagens raster (Pillow)
# ---------------------------------------------------------------------------

_HEIF_DONE = False


def _register_heif():
    global _HEIF_DONE
    if _HEIF_DONE:
        return
    try:
        import pillow_heif
        pillow_heif.register_heif_opener()
    except Exception:
        pass
    _HEIF_DONE = True


PIL_FORMAT = {
    'png': 'PNG', 'jpg': 'JPEG', 'webp': 'WEBP', 'bmp': 'BMP', 'gif': 'GIF', 'tiff': 'TIFF',
    'ico': 'ICO', 'icns': 'ICNS', 'heic': 'HEIF', 'avif': 'AVIF', 'tga': 'TGA', 'ppm': 'PPM',
    'pcx': 'PCX', 'jp2': 'JPEG2000', 'qoi': 'QOI', 'pdf': 'PDF',
}
NO_ALPHA = {'jpg', 'ppm', 'pcx', 'pdf', 'eps'}
ANIMATED_OUT = {'gif', 'webp', 'png', 'tiff', 'avif'}


def _open_image(path):
    from PIL import Image, ImageOps
    _register_heif()
    try:
        img = Image.open(path)
        img.load()
    except Exception as e:
        raise ConversionError(f"Não foi possível abrir a imagem ({e}).")
    return img, ImageOps


def _flatten(img, bg=(255, 255, 255)):
    from PIL import Image
    if img.mode in ('RGBA', 'LA') or (img.mode == 'P' and 'transparency' in img.info):
        rgba = img.convert('RGBA')
        base = Image.new('RGB', rgba.size, bg)
        base.paste(rgba, mask=rgba.split()[-1])
        return base
    if img.mode not in ('RGB', 'L'):
        return img.convert('RGB')
    return img


def _prepare_frame(frame, out, opts):
    from PIL import Image
    width = _opt_int(opts, 'resize_width', 0, 0, 20000)
    if width and width != frame.width:
        h = max(1, round(frame.height * width / frame.width))
        frame = frame.resize((width, h), Image.LANCZOS)
    if out in NO_ALPHA:
        frame = _flatten(frame)
    elif out == 'gif':
        return frame
    elif frame.mode not in ('RGB', 'RGBA', 'L', 'LA'):
        frame = frame.convert('RGBA' if 'transparency' in frame.info or frame.mode in ('P', 'PA') else 'RGB')
    if out in ('bmp', 'tga', 'jp2', 'qoi', 'avif', 'heic', 'webp', 'ico', 'icns') and frame.mode in ('LA', 'P'):
        frame = frame.convert('RGBA')
    if out == 'ppm' and frame.mode not in ('RGB', 'L'):
        frame = frame.convert('RGB')
    return frame


def convert_raster(src, out, opts, work, base):
    img, ImageOps = _open_image(src)
    quality = _opt_int(opts, 'quality', 90, 1, 100)
    out_path = os.path.join(work, f"{base}.{out}")
    n_frames = getattr(img, 'n_frames', 1)
    save_kwargs = {}

    if out in ('jpg', 'webp', 'avif', 'heic', 'jp2'):
        save_kwargs['quality'] = quality
        if out == 'jpg':
            save_kwargs.update(optimize=True, progressive=True)
            if img.info.get('icc_profile'):
                save_kwargs['icc_profile'] = img.info['icc_profile']
        if out == 'jp2':
            save_kwargs = {'quality_mode': 'dB', 'quality_layers': [30 + quality * 0.2]}
    if out == 'png':
        save_kwargs['optimize'] = True
    if out == 'tiff':
        save_kwargs['compression'] = 'tiff_lzw'
    if out == 'ico':
        size = max(img.size)
        save_kwargs['sizes'] = [(s, s) for s in (16, 24, 32, 48, 64, 128, 256) if s <= max(size, 16)]
    if out == 'pdf':
        save_kwargs['resolution'] = float((img.info.get('dpi') or (150, 150))[0] or 150)

    if out == 'heic':
        _register_heif()

    try:
        if n_frames > 1 and out in ANIMATED_OUT:
            from PIL import ImageSequence
            frames = [_prepare_frame(f.copy(), out, opts) for f in ImageSequence.Iterator(img)]
            durations = []
            for f in ImageSequence.Iterator(img):
                durations.append(f.info.get('duration', img.info.get('duration', 100)))
            # Todos os quadros precisam do mesmo modo (o GIF mistura P/RGB/RGBA entre quadros)
            if out == 'gif':
                frames = [f.convert('RGBA') if f.mode not in ('P', 'L') else f for f in frames]
            elif out != 'tiff' or len({f.mode for f in frames}) > 1:
                frames = [f.convert('RGBA') for f in frames]
            extra = {'save_all': True, 'append_images': frames[1:], 'loop': img.info.get('loop', 0)}
            if out in ('gif', 'webp', 'png', 'avif'):
                extra['duration'] = durations
            if out == 'gif':
                extra['disposal'] = 2
            frames[0].save(out_path, format=PIL_FORMAT[out], **save_kwargs, **extra)
        else:
            frame = ImageOps.exif_transpose(img)
            frame = _prepare_frame(frame, out, opts)
            if out == 'icns':
                side = max(frame.size)
                square = frame.convert('RGBA').resize((side, side)) if frame.width != frame.height else frame
                frame = square.resize((1024, 1024)) if side < 1024 else square
            if out == 'pdf' and frame.mode not in ('RGB', 'L'):
                frame = frame.convert('RGB')
            frame.save(out_path, format=PIL_FORMAT[out], **save_kwargs)
    except (KeyError, OSError, ValueError, TypeError) as e:
        raise ConversionError(f"O Pillow não conseguiu gerar {out.upper()}: {e}")
    return [out_path]


# ---------------------------------------------------------------------------
# Vetorização (imagem -> DXF / SVG / EPS)
# ---------------------------------------------------------------------------

class TraceResult:
    """Resultado da vetorização em coordenadas de pixel da imagem original
    (origem no canto superior esquerdo, y para baixo).

    layers: lista de dicts {name, color(rgb), kind('fill'|'line'), curves}
      - curves (kind fill): lista de subcaminhos fechados, cada um uma lista de
        comandos ('M', (x,y)) / ('L', (x,y)) / ('C', c1, c2, end)
      - curves (kind line): lista de polilinhas abertas/fechadas [(x,y), ...]
    """

    def __init__(self, width, height, dpi, layers):
        self.width = width
        self.height = height
        self.dpi = dpi
        self.layers = layers

    def count(self):
        return sum(len(l['curves']) for l in self.layers)


def _load_trace_source(src, opts):
    """Abre a imagem, lida com transparência e devolve (rgb, alpha|None, dpi, escala)."""
    from PIL import Image
    img, ImageOps = _open_image(src)
    img = ImageOps.exif_transpose(img)
    dpi = img.info.get('dpi', (96, 96))
    try:
        dpi = float(dpi[0]) if dpi and dpi[0] else 96.0
    except (TypeError, ValueError, IndexError):
        dpi = 96.0
    if dpi < 10 or dpi > 4800:
        dpi = 96.0

    has_alpha = img.mode in ('RGBA', 'LA', 'PA') or (img.mode == 'P' and 'transparency' in img.info)
    rgba = img.convert('RGBA') if has_alpha else img.convert('RGB')
    orig_w, orig_h = rgba.size

    # Trabalha numa resolução confortável para o potrace: imagens pequenas são
    # ampliadas (curvas bem mais suaves, como fazem os conversores online) e as
    # enormes são reduzidas (o ganho de detalhe não compensa o tempo).
    longest = max(orig_w, orig_h)
    scale = 1.0
    max_side = TRACE_MAX_SIDE_COLOR if opts.get('trace_mode') == 'color' else TRACE_MAX_SIDE
    if opts.get('_max_side'):
        max_side = min(max_side, int(opts['_max_side']))
    if longest > max_side:
        scale = max_side / longest
    elif longest < TRACE_MIN_SIDE:
        scale = min(4.0, TRACE_MIN_SIDE / longest)
    if scale != 1.0:
        new_size = (max(1, round(orig_w * scale)), max(1, round(orig_h * scale)))
        rgba = rgba.resize(new_size, Image.LANCZOS if scale < 1 else Image.BICUBIC)

    alpha = None
    if has_alpha:
        alpha = rgba.split()[-1]
        rgb = Image.new('RGB', rgba.size, (255, 255, 255))
        rgb.paste(rgba, mask=alpha)
    else:
        rgb = rgba
    return rgb, alpha, dpi, scale, (orig_w, orig_h)


def _ink_mask(rgb, alpha, opts):
    """Máscara booleana True = "tinta" (o que vira vetor)."""
    import numpy as np
    gray = np.asarray(rgb.convert('L'), dtype=np.uint8)
    thr_raw = str(opts.get('threshold', 'auto')).strip().lower()
    if thr_raw in ('', 'auto'):
        try:
            from skimage.filters import threshold_otsu
            thr = float(threshold_otsu(gray)) if gray.min() != gray.max() else 128.0
        except Exception:
            thr = 128.0
    else:
        thr = float(_opt_int(opts, 'threshold', 128, 1, 254))
    ink = gray < thr

    if alpha is not None:
        a = np.asarray(alpha, dtype=np.uint8)
        transparent = a < 128
        # Logo em fundo transparente: se o que é opaco é quase tudo "claro" (ex.:
        # logo branco), a forma é definida pela opacidade e não pela cor.
        opaque = ~transparent
        if transparent.mean() > 0.01 and opaque.any():
            if ink[opaque].mean() < 0.05:
                ink = opaque
            else:
                ink &= opaque

    if _opt_bool(opts, 'invert'):
        ink = ~ink
    return ink


def _potrace_curves(ink, opts):
    """Roda o potrace numa máscara e devolve subcaminhos no formato de TraceResult."""
    import potrace
    import numpy as np
    if not ink.any():
        return []
    turdsize = _opt_int(opts, 'despeckle', 2, 0, 500)
    alphamax = _opt_float(opts, 'smooth', 1.0, 0.0, 1.334)
    opttol = _opt_float(opts, 'optimize', 0.2, 0.0, 5.0)
    bitmap = potrace.Bitmap(~np.asarray(ink, dtype=bool))  # potracer: True = fundo
    plist = bitmap.trace(turdsize=turdsize, turnpolicy=potrace.POTRACE_TURNPOLICY_MINORITY,
                         alphamax=alphamax, opticurve=opttol > 0, opttolerance=opttol or 0.2)
    subpaths = []
    for curve in plist:
        sp = curve.start_point
        cmds = [('M', (sp.x, sp.y))]
        for seg in curve.segments:
            if seg.is_corner:
                cmds.append(('L', (seg.c.x, seg.c.y)))
                cmds.append(('L', (seg.end_point.x, seg.end_point.y)))
            else:
                cmds.append(('C', (seg.c1.x, seg.c1.y), (seg.c2.x, seg.c2.y),
                             (seg.end_point.x, seg.end_point.y)))
        subpaths.append(cmds)
    return subpaths


def _centerline(ink, opts):
    """Esqueletiza o traço e segue os pixels do esqueleto, gerando polilinhas
    abertas no centro de cada linha (uma passada só da ferramenta no CNC/laser)."""
    import numpy as np
    from skimage.morphology import skeletonize, remove_small_objects
    from skimage.measure import approximate_polygon

    min_px = _opt_int(opts, 'despeckle', 2, 0, 500)
    mask = np.asarray(ink, dtype=bool)
    if min_px:
        mask = remove_small_objects(mask, max_size=max(1, min_px * 4))
    skel = skeletonize(mask)
    if not skel.any():
        return []

    h, w = skel.shape
    padded = np.pad(skel, 1)
    ys, xs = np.nonzero(skel)
    offsets = [(-1, -1), (-1, 0), (-1, 1), (0, -1), (0, 1), (1, -1), (1, 0), (1, 1)]
    pixels = set(zip(ys.tolist(), xs.tolist()))

    def neighbors(p):
        y, x = p
        out = []
        for dy, dx in offsets:
            q = (y + dy, x + dx)
            if padded[q[0] + 1, q[1] + 1]:
                out.append(q)
        return out

    degree = {p: len(neighbors(p)) for p in pixels}
    nodes = {p for p, d in degree.items() if d != 2}
    visited_edges = set()

    def edge_key(a, b):
        return (a, b) if a <= b else (b, a)

    polylines = []

    def walk(start, nxt):
        line = [start, nxt]
        visited_edges.add(edge_key(start, nxt))
        prev, cur = start, nxt
        while cur not in nodes:
            step = None
            for q in neighbors(cur):
                if q != prev and edge_key(cur, q) not in visited_edges:
                    step = q
                    break
            if step is None:
                break
            visited_edges.add(edge_key(cur, step))
            line.append(step)
            prev, cur = cur, step
            if cur == start:
                break
        return line

    for node in nodes:
        for q in neighbors(node):
            if edge_key(node, q) not in visited_edges:
                polylines.append(walk(node, q))
    # Laços fechados sem nenhuma ponta/junção (ex.: um "O")
    for p in pixels:
        for q in neighbors(p):
            if edge_key(p, q) not in visited_edges:
                polylines.append(walk(p, q))

    # Espessura típica do traço (2x a distância do esqueleto até a borda)
    from scipy.ndimage import distance_transform_edt
    dist = distance_transform_edt(mask)
    stroke = float(np.median(dist[skel])) * 2 if skel.any() else 2.0

    def plen(line):
        a = np.array(line, dtype=float)
        return float(np.sum(np.hypot(*np.diff(a, axis=0).T))) if len(a) > 1 else 0.0

    # Esporões: galhos curtos que saem de uma junção e terminam soltos são
    # artefatos do esqueleto (aparecem em cantos e cruzamentos), não desenho
    def is_spur(line):
        a, b = line[0], line[-1]
        da, db = degree.get(a, 2), degree.get(b, 2)
        dangling = (da == 1 and db >= 3) or (db == 1 and da >= 3)
        return dangling and plen(line) < stroke * 1.5 + 2

    polylines = [l for l in polylines if len(l) >= 2 and not is_spur(l)]

    # Junta linhas que ficaram partidas numa junção onde só sobraram duas
    changed = True
    while changed:
        changed = False
        ends = {}
        for i, l in enumerate(polylines):
            if l[0] == l[-1]:
                continue
            for end in (l[0], l[-1]):
                ends.setdefault(end, []).append(i)
        for node, idxs in ends.items():
            if len(idxs) == 2 and idxs[0] != idxs[1] and degree.get(node, 2) >= 3:
                a, b = polylines[idxs[0]], polylines[idxs[1]]
                if a[-1] != node:
                    a = a[::-1]
                if b[0] != node:
                    b = b[::-1]
                merged = a + b[1:]
                polylines = [l for j, l in enumerate(polylines) if j not in idxs] + [merged]
                changed = True
                break

    tolerance = max(0.6, _opt_float(opts, 'optimize', 0.2, 0.0, 5.0) * 3)
    min_len = max(2, min_px)
    result = []
    for line in polylines:
        coords = np.array([(x + 0.5, y + 0.5) for y, x in line], dtype=float)
        if plen(line) < min_len:
            continue
        simplified = approximate_polygon(coords, tolerance=tolerance)
        result.append([tuple(pt) for pt in simplified.tolist()])
    return result


def _scale_cmds(subpaths, k):
    if k == 1.0:
        return subpaths
    out = []
    for cmds in subpaths:
        new = []
        for c in cmds:
            new.append((c[0],) + tuple((p[0] * k, p[1] * k) for p in c[1:]))
        out.append(new)
    return out


def trace_image(src, opts):
    """Vetoriza a imagem conforme opts['trace_mode']: outline | centerline | color."""
    import numpy as np
    mode = opts.get('trace_mode', 'outline')
    rgb, alpha, dpi, scale, (orig_w, orig_h) = _load_trace_source(src, opts)
    k = 1.0 / scale
    layers = []

    if mode == 'color':
        from PIL import Image, ImageFilter
        n_colors = _opt_int(opts, 'colors', 6, 2, 32)
        smoothed = rgb.filter(ImageFilter.MedianFilter(3))
        transparent = np.asarray(alpha) < 128 if alpha is not None else None
        # A paleta sai só de pixels visíveis e "lisos" (vizinhança uniforme):
        # a área transparente (que vira branco) e as faixas de antisserrilhado
        # nas bordas roubariam cores pedidas e virariam camadas falsas. Depois,
        # todo pixel (inclusive os de borda) vai para a cor mais próxima.
        from scipy.ndimage import maximum_filter, minimum_filter
        arr = np.asarray(smoothed)
        gray = np.asarray(smoothed.convert('L'))
        flat = (maximum_filter(gray, size=5).astype(int) - minimum_filter(gray, size=5)) < 16
        sample = flat if transparent is None else (flat & ~transparent)
        if sample.sum() < 500:
            sample = np.ones(gray.shape, bool) if transparent is None else ~transparent
        if sample.any():
            pal_img = Image.fromarray(arr[sample].reshape(1, -1, 3)).quantize(
                colors=n_colors, method=Image.Quantize.MEDIANCUT)
            quant = smoothed.quantize(palette=pal_img, dither=Image.Dither.NONE)
        else:
            quant = smoothed.quantize(colors=n_colors, method=Image.Quantize.MEDIANCUT,
                                      dither=Image.Dither.NONE)
        idx = np.asarray(quant).copy()
        palette = quant.getpalette()[:n_colors * 3]
        colors = [tuple(palette[i * 3:i * 3 + 3]) for i in range(len(palette) // 3)]

        # O antisserrilhado das bordas cria "cores de transição" em faixas finas;
        # cores com menos de 0,4% da imagem são fundidas na cor principal mais
        # próxima, senão cada borda viraria uma camada extra cheia de recortes.
        visible = idx[~transparent] if transparent is not None else idx.ravel()
        counts = np.bincount(visible, minlength=len(colors))
        order = [int(i) for i in np.argsort(-counts)]
        pal = np.array(colors, dtype=float)
        major = []
        for i in order:
            if counts[i] < max(1, visible.size * 0.004):
                continue
            # Tons quase iguais (o median cut às vezes divide uma cor lisa em duas)
            if any(((pal[m] - pal[i]) ** 2).sum() < 20 ** 2 for m in major):
                continue
            major.append(i)
        major = major or [order[0]]
        remap = np.arange(len(colors))
        for i in range(len(colors)):
            if i not in major:
                remap[i] = major[int(np.argmin(((pal[major] - pal[i]) ** 2).sum(axis=1)))]
        idx = remap[idx].astype(np.uint8)
        # Filtro de moda: cada pixel assume a cor predominante da vizinhança,
        # removendo serrilhado e pixels soltos nas fronteiras entre cores
        idx = np.asarray(Image.fromarray(idx, 'L').filter(ImageFilter.ModeFilter(5)))

        # Fundo = cor mais frequente na borda. Só é descartado se for branco (ou
        # transparente): numa foto o "fundo" é parte da imagem e sumir com ele
        # deixaria buracos.
        if transparent is not None and transparent.mean() > 0.01:
            bg_index = -1
        else:
            border = np.concatenate([idx[0, :], idx[-1, :], idx[:, 0], idx[:, -1]])
            bg_index = int(np.bincount(border).argmax())
            if min(colors[bg_index]) < 225:
                bg_index = -1
        if _opt_bool(opts, 'keep_background'):
            bg_index = -1

        entries = []
        for ci, color in enumerate(colors):
            mask = idx == ci
            if transparent is not None:
                mask &= ~transparent
            if not mask.any() or ci == bg_index:
                continue
            entries.append((int(mask.sum()), ci, color, mask))
        entries.sort(key=lambda e: -e[0])

        stacked = opts.get('_stacked', True)
        for pos, (_, ci, color, mask) in enumerate(entries):
            if stacked:
                # "Empilhado": cada camada cobre também as camadas desenhadas por
                # cima dela, eliminando frestas brancas entre cores vizinhas.
                full = mask.copy()
                for _, _, _, above in entries[pos + 1:]:
                    full |= above
                mask = full
            curves = _scale_cmds(_potrace_curves(mask, opts), k)
            if curves:
                layers.append({'name': f"COR_{pos + 1}_{_hex(color)[1:].upper()}",
                               'color': color, 'kind': 'fill', 'curves': curves})
    elif mode == 'centerline':
        ink = _ink_mask(rgb, alpha, opts)
        lines = _centerline(ink, opts)
        lines = [[(x * k, y * k) for x, y in line] for line in lines]
        if lines:
            layers.append({'name': 'LINHA_CENTRAL', 'color': (0, 0, 0), 'kind': 'line', 'curves': lines})
    else:
        ink = _ink_mask(rgb, alpha, opts)
        curves = _scale_cmds(_potrace_curves(ink, opts), k)
        if curves:
            layers.append({'name': 'CONTORNO', 'color': (0, 0, 0), 'kind': 'fill', 'curves': curves})

    result = TraceResult(orig_w, orig_h, dpi, layers)
    if not result.count():
        raise ConversionError(
            "Nenhuma forma foi encontrada na imagem. Ajuste o limiar (ou marque "
            "'Inverter') - a imagem pode estar clara/escura demais para o limiar atual.")
    return result


def _fmt(v):
    s = f"{v:.2f}".rstrip('0').rstrip('.')
    return s if s not in ('-0', '') else '0'


def trace_to_svg(tr, stroke_fill_gap=True):
    parts = [
        '<?xml version="1.0" encoding="UTF-8"?>',
        f'<svg xmlns="http://www.w3.org/2000/svg" version="1.1" width="{tr.width}" '
        f'height="{tr.height}" viewBox="0 0 {tr.width} {tr.height}">',
    ]
    for layer in tr.layers:
        color = _hex(layer['color'])
        if layer['kind'] == 'fill':
            d = []
            for cmds in layer['curves']:
                for c in cmds:
                    if c[0] == 'M':
                        d.append(f"M{_fmt(c[1][0])} {_fmt(c[1][1])}")
                    elif c[0] == 'L':
                        d.append(f"L{_fmt(c[1][0])} {_fmt(c[1][1])}")
                    else:
                        d.append("C" + " ".join(f"{_fmt(p[0])} {_fmt(p[1])}" for p in c[1:]))
                d.append('Z')
            parts.append(f'<path id="{layer["name"]}" fill="{color}" fill-rule="evenodd" d="{"".join(d)}"/>')
        else:
            sw = max(1.0, max(tr.width, tr.height) / 800)
            parts.append(f'<g id="{layer["name"]}" fill="none" stroke="{color}" stroke-width="{_fmt(sw)}" '
                         f'stroke-linecap="round" stroke-linejoin="round">')
            for line in layer['curves']:
                pts = " ".join(f"{_fmt(x)},{_fmt(y)}" for x, y in line)
                parts.append(f'<polyline points="{pts}"/>')
            parts.append('</g>')
    parts.append('</svg>')
    return "\n".join(parts)


def trace_to_eps(tr):
    """EPS vetorial escrito à mão (PostScript é simples: moveto/curveto/eofill)."""
    k = 72.0 / tr.dpi
    W, H = tr.width * k, tr.height * k
    out = [
        '%!PS-Adobe-3.0 EPSF-3.0',
        f'%%BoundingBox: 0 0 {math.ceil(W)} {math.ceil(H)}',
        f'%%HiResBoundingBox: 0 0 {W:.3f} {H:.3f}',
        '%%Creator: Verto Conversor Universal',
        '%%EndComments',
        'gsave',
    ]

    def P(p):
        return f"{p[0] * k:.3f} {(tr.height - p[1]) * k:.3f}"

    for layer in tr.layers:
        r, g, b = (c / 255.0 for c in layer['color'])
        out.append(f"{r:.4f} {g:.4f} {b:.4f} setrgbcolor")
        if layer['kind'] == 'fill':
            out.append('newpath')
            for cmds in layer['curves']:
                for c in cmds:
                    if c[0] == 'M':
                        out.append(f"{P(c[1])} moveto")
                    elif c[0] == 'L':
                        out.append(f"{P(c[1])} lineto")
                    else:
                        out.append(f"{P(c[1])} {P(c[2])} {P(c[3])} curveto")
                out.append('closepath')
            out.append('eofill')
        else:
            out.append(f"{max(0.3, 72.0 / tr.dpi):.3f} setlinewidth 1 setlinecap 1 setlinejoin")
            for line in layer['curves']:
                out.append('newpath')
                out.append(f"{P(line[0])} moveto")
                for p in line[1:]:
                    out.append(f"{P(p)} lineto")
                out.append('stroke')
    out += ['grestore', 'showpage', '%%EOF']
    return "\n".join(out) + "\n"


DXF_VERSIONS = {'R12': 'R12', 'R2000': 'R2000', 'R2004': 'R2004', 'R2007': 'R2007',
                'R2010': 'R2010', 'R2013': 'R2013', 'R2018': 'R2018'}


def _new_dxf(opts):
    import ezdxf
    from ezdxf import units
    version = DXF_VERSIONS.get(str(opts.get('dxf_version', 'R2010')).upper(), 'R2010')
    # R12 é gerado a partir de um R2010 e exportado no fim (ezdxf.addons.r12export)
    doc = ezdxf.new('R2010' if version == 'R12' else version, setup=True)
    doc.units = units.MM
    doc.header['$MEASUREMENT'] = 1
    doc.header['$INSUNITS'] = 4
    return doc, version


def _finalize_dxf(doc):
    """Grava no cabeçalho a extensão real do desenho ($EXTMIN/$EXTMAX, $LIMMIN/
    $LIMMAX) e deixa a janela inicial enquadrando tudo. Sem isso o ezdxf deixa
    os valores "vazios" (±1e20) e vários programas (LibreOffice, visualizadores,
    alguns CAM) abrem o desenho minúsculo, num ponto só ou fora da tela."""
    from ezdxf import bbox
    msp = doc.modelspace()
    ext = bbox.extents(msp, fast=True)
    if not ext.has_data:
        return
    (x0, y0, z0), (x1, y1, z1) = ext.extmin, ext.extmax
    # O ezdxf copia a extensão do layout para o cabeçalho ao salvar, mas pula a
    # cópia quando extmin é (0, 0, 0) (Vec3 nulo é "falso") - caso comum, pois a
    # imagem vetorizada começa na origem. Por isso grava nos dois lugares.
    msp.reset_extents((x0, y0, z0), (x1, y1, z1))
    msp.reset_limits((x0, y0), (x1, y1))
    doc.header['$EXTMIN'] = (x0, y0, z0)
    doc.header['$EXTMAX'] = (x1, y1, z1)
    w, h = max(x1 - x0, 1e-6), max(y1 - y0, 1e-6)
    try:
        doc.set_modelspace_vport(height=max(h, w * 0.75) * 1.1, center=((x0 + x1) / 2, (y0 + y1) / 2))
    except Exception:
        pass


def _save_dxf(doc, version, out_path):
    _finalize_dxf(doc)
    if version == 'R12':
        from ezdxf.addons import r12export
        r12export.saveas(doc, out_path)
    else:
        doc.saveas(out_path)


def _dxf_true_color(rgb):
    from ezdxf import colors
    return colors.rgb2int(tuple(int(c) for c in rgb[:3]))


def _render_paths(msp, paths, opts, attribs):
    """Escreve caminhos ezdxf.path no modelspace como polilinha, spline ou hachura."""
    from ezdxf import path as ezpath
    entity = opts.get('dxf_entity', 'polyline')
    if not paths:
        return
    tol = _opt_float(opts, 'dxf_tolerance', 0.05, 0.001, 5.0)
    if entity == 'spline':
        ezpath.render_splines_and_polylines(msp, paths, dxfattribs=attribs)
    else:
        ezpath.render_lwpolylines(msp, paths, distance=tol, segments=4, dxfattribs=attribs)
    if _opt_bool(opts, 'dxf_hatch'):
        hatch_attr = dict(attribs)
        hatch_attr['layer'] = attribs.get('layer', '0') + '_PREENCHIMENTO'
        layers = msp.doc.layers
        if hatch_attr['layer'] not in layers:
            layer = layers.add(hatch_attr['layer'], color=7)
            if 'true_color' in attribs:
                layer.dxf.true_color = attribs['true_color']
        _render_nested_hatches(msp, paths, tol, hatch_attr)


def _nest_polygons(polys):
    """Árvore de encaixe: para cada polígono, o índice do menor polígono que o
    contém (ou -1) e a profundidade. Profundidade par = forma, ímpar = furo."""
    import numpy as np
    n = len(polys)
    boxes = np.array([(p[:, 0].min(), p[:, 1].min(), p[:, 0].max(), p[:, 1].max()) for p in polys])
    areas = np.array([abs(0.5 * np.sum(p[:, 0] * np.roll(p[:, 1], -1) - np.roll(p[:, 0], -1) * p[:, 1]))
                      for p in polys])

    def inside(pt, poly):
        x, y = pt
        xs, ys = poly[:, 0], poly[:, 1]
        xj, yj = np.roll(xs, 1), np.roll(ys, 1)
        cross = ((ys > y) != (yj > y)) & (x < (xj - xs) * (y - ys) / np.where(yj - ys == 0, 1e-12, yj - ys) + xs)
        return bool(np.count_nonzero(cross) % 2)

    parent = [-1] * n
    for i in range(n):
        pt = polys[i][0]
        b = boxes[i]
        cand = np.nonzero((boxes[:, 0] <= b[0]) & (boxes[:, 1] <= b[1]) &
                          (boxes[:, 2] >= b[2]) & (boxes[:, 3] >= b[3]) & (areas > areas[i]))[0]
        best, best_area = -1, None
        for j in cand:
            if j != i and inside(pt, polys[j]) and (best_area is None or areas[j] < best_area):
                best, best_area = int(j), areas[j]
        parent[i] = best
    depth = [0] * n
    for i in range(n):
        d, j = 0, parent[i]
        while j != -1 and d < n:
            d += 1
            j = parent[j]
        depth[i] = d
    return parent, depth


def _orient(a, b, c):
    return (b[..., 0] - a[..., 0]) * (c[..., 1] - a[..., 1]) - (b[..., 1] - a[..., 1]) * (c[..., 0] - a[..., 0])


def _segment_clear(ring, m, p):
    """True se o segmento m-p não cruza nenhuma aresta do anel (toques nas
    pontas não contam, pois a ponte começa/termina em vértices)."""
    import numpy as np
    a, b = ring, np.roll(ring, -1, axis=0)
    m_ = np.broadcast_to(m, a.shape)
    p_ = np.broadcast_to(p, a.shape)
    d1, d2 = _orient(m_, p_, a), _orient(m_, p_, b)
    d3, d4 = _orient(a, b, m_), _orient(a, b, p_)
    eps = 1e-12
    cross = (((d1 > eps) & (d2 < -eps)) | ((d1 < -eps) & (d2 > eps))) & \
            (((d3 > eps) & (d4 < -eps)) | ((d3 < -eps) & (d4 > eps)))
    return not cross.any()


def _bridge_holes(outer, holes):
    """Une os furos ao contorno externo com "pontes" de largura zero (keyhole),
    como faz o earcut: devolve UM laço simples com a mesma área pintada.
    Visualizadores que ignoram furos de hachura (ou pintam cada laço como forma
    cheia) passam a mostrar o desenho certo, porque não existe mais furo.

    outer anti-horário, furos horários. Os furos são processados da esquerda
    para a direita; a ponte sai do ponto mais à esquerda do furo em direção a
    um vértice visível do anel já montado (que contém os furos anteriores)."""
    import numpy as np
    ring = np.asarray(outer, dtype=float)
    order = sorted(range(len(holes)), key=lambda i: float(np.min(holes[i][:, 0])))
    for i in order:
        h = np.asarray(holes[i], dtype=float)
        k = int(np.lexsort((h[:, 1], h[:, 0]))[0])  # mais à esquerda (e mais baixo no empate)
        mx, my = h[k]
        m = h[k]
        a, b = ring, np.roll(ring, -1, axis=0)
        crosses = (a[:, 1] > my) != (b[:, 1] > my)
        with np.errstate(divide='ignore', invalid='ignore'):
            xi = a[:, 0] + (my - a[:, 1]) * (b[:, 0] - a[:, 0]) / (b[:, 1] - a[:, 1])
        valid = crosses & (xi <= mx)
        target = None
        if valid.any():
            j = int(np.flatnonzero(valid)[np.argmax(xi[valid])])
            qx = xi[j]
            jn = (j + 1) % len(ring)
            cand = j if ring[j, 0] < ring[jn, 0] else jn
            # Vértices do anel dentro do triângulo (furo, interseção, candidato)
            # bloqueariam a ponte: escolhe o de menor ângulo com o raio
            px, py = ring[cand]
            tri = np.array([[mx, my], [qx, my], [px, py]])
            pts = ring
            sel = (pts[:, 0] <= mx) & (pts[:, 0] >= px) & (pts[:, 0] != mx)
            if sel.any():
                t = np.broadcast_to
                s1 = _orient(t(tri[0], pts.shape), t(tri[1], pts.shape), pts)
                s2 = _orient(t(tri[1], pts.shape), t(tri[2], pts.shape), pts)
                s3 = _orient(t(tri[2], pts.shape), t(tri[0], pts.shape), pts)
                inside = sel & (((s1 >= 0) & (s2 >= 0) & (s3 >= 0)) | ((s1 <= 0) & (s2 <= 0) & (s3 <= 0)))
                inside[cand] = False
                if inside.any():
                    idx = np.flatnonzero(inside)
                    tan = np.abs(my - pts[idx, 1]) / np.maximum(mx - pts[idx, 0], 1e-12)
                    cand = int(idx[np.argmin(tan)])
            if _segment_clear(ring, m, ring[cand]):
                target = cand
        if target is None:
            # Plano B: o vértice visível mais próximo à esquerda do furo
            dist = np.hypot(ring[:, 0] - mx, ring[:, 1] - my)
            dist[ring[:, 0] > mx] = np.inf
            for c in np.argsort(dist)[:400]:
                if not np.isfinite(dist[c]):
                    break
                if _segment_clear(ring, m, ring[c]):
                    target = int(c)
                    break
        if target is None:
            return None  # não achou ponte segura: quem chamou usa o formato com furos
        hole_loop = np.vstack([h[k:], h[:k], h[k:k + 1]])  # começa e termina em m
        ring = np.vstack([ring[:target + 1], hole_loop, ring[target:target + 1], ring[target + 1:]])
    return ring


def _render_nested_hatches(msp, paths, tol, attribs):
    """Uma HATCH sólida por forma: contorno externo + furos diretos. Ilhas dentro
    de furos viram hachuras próprias. É a estrutura que todo programa CAD
    entende - uma hachura única com centenas de contornos depende da regra
    par-ímpar, e vários programas (LibreOffice, alguns visualizadores e CAMs)
    acabam pintando os furos."""
    import numpy as np
    from ezdxf import const
    polys = []
    for p in paths:
        if not p.is_closed:
            continue  # linha solta (comum em PDF/SVG) não tem área para preencher
        pts = [(v.x, v.y) for v in p.flattening(distance=tol, segments=4)]
        if len(pts) > 2 and pts[0] == pts[-1]:
            pts = pts[:-1]
        if len(pts) >= 3:
            polys.append(np.array(pts, dtype=float))
    if not polys:
        return
    parent, depth = _nest_polygons(polys)
    holes = {}
    for i, par in enumerate(parent):
        if depth[i] % 2 == 1 and par != -1:
            holes.setdefault(par, []).append(i)
    layers = msp.doc.layers
    if attribs.get('layer') and attribs['layer'] not in layers:
        layers.add(attribs['layer'], color=7)
    def oriented(poly, ccw):
        # Externo anti-horário e furos horários: assim a hachura sai certa tanto
        # em programas que usam a regra par-ímpar quanto nos que usam o sentido
        # de giro (nonzero), onde um furo "no mesmo sentido" seria pintado
        area2 = np.sum(poly[:, 0] * np.roll(poly[:, 1], -1) - np.roll(poly[:, 0], -1) * poly[:, 1])
        return (poly if (area2 > 0) == ccw else poly[::-1]).tolist()

    outer_flags = const.BOUNDARY_PATH_EXTERNAL
    for i, poly in enumerate(polys):
        if depth[i] % 2 == 1:
            continue
        hatch = msp.add_hatch(dxfattribs=dict(attribs))
        # Estilo Normal (par-ímpar, o padrão do HATCH do AutoCAD): usa todos os
        # contornos. Os furos vão marcados como "outermost" para também valerem
        # em leitores que aplicam o estilo "Externo", que descarta os demais.
        hatch.set_solid_fill(color=256, style=const.HATCH_STYLE_NESTED)  # cor da camada
        if 'true_color' in attribs:
            hatch.dxf.true_color = attribs['true_color']
        outer = np.array(oriented(poly, True))
        inner = [np.array(oriented(polys[h], False)) for h in holes.get(i, [])]
        merged = _bridge_holes(outer, inner) if inner else outer
        if merged is not None:
            # Um laço só, sem furos: igual em qualquer visualizador
            hatch.paths.add_polyline_path(merged.tolist(), is_closed=True, flags=outer_flags)
        else:
            hatch.paths.add_polyline_path(outer.tolist(), is_closed=True, flags=outer_flags)
            for hole in inner:
                hatch.paths.add_polyline_path(hole.tolist(), is_closed=True,
                                              flags=const.BOUNDARY_PATH_OUTERMOST)


def trace_to_dxf(tr, opts, out_path):
    from ezdxf import path as ezpath
    from ezdxf.math import Vec2

    doc, version = _new_dxf(opts)
    msp = doc.modelspace()

    # Escala: largura final em mm (se informada) ou tamanho físico pelo DPI da imagem
    width_mm = _opt_float(opts, 'width_mm', 0.0, 0.0, 100000.0)
    k = (width_mm / tr.width) if width_mm > 0 else (25.4 / tr.dpi)
    H = tr.height

    def V(p):
        return Vec2(p[0] * k, (H - p[1]) * k)

    for layer in tr.layers:
        name = layer['name']
        if name not in doc.layers:
            doc.layers.add(name, color=7 if layer['color'] == (0, 0, 0) else 256)
            if layer['color'] != (0, 0, 0):
                doc.layers.get(name).dxf.true_color = _dxf_true_color(layer['color'])
        attribs = {'layer': name}
        if layer['color'] != (0, 0, 0):
            attribs['true_color'] = _dxf_true_color(layer['color'])

        if layer['kind'] == 'line':
            for line in layer['curves']:
                pts = [V(p) for p in line]
                closed = len(pts) > 2 and pts[0].isclose(pts[-1], abs_tol=1e-6)
                if closed:
                    pts = pts[:-1]
                msp.add_lwpolyline(pts, close=closed, dxfattribs=attribs)
            continue

        paths = []
        for cmds in layer['curves']:
            p = None
            for c in cmds:
                if c[0] == 'M':
                    p = ezpath.Path(V(c[1]))
                elif c[0] == 'L':
                    if not V(c[1]).isclose(p.end, abs_tol=1e-9):
                        p.line_to(V(c[1]))
                else:
                    p.curve4_to(V(c[3]), V(c[1]), V(c[2]))
            if p is not None:
                p.close()
                paths.append(p)
        _render_paths(msp, paths, opts, attribs)

    _save_dxf(doc, version, out_path)


def convert_trace(src, out, opts, work, base):
    opts = dict(opts)
    # Preenchido por padrão, como na prévia: só com os contornos, uma foto
    # vetorizada vira um "desenho de bordas" que não lembra a imagem
    if out == 'dxf' and opts.get('trace_mode') != 'centerline' and opts.get('dxf_hatch') in (None, ''):
        opts['dxf_hatch'] = '1'
    # No DXF as cores ficam lado a lado (sem sobreposição) - é o que CNC/laser
    # espera; no SVG/EPS as camadas se sobrepõem para não sobrar fresta branca.
    opts['_stacked'] = out != 'dxf'
    tr = trace_image(src, opts)
    out_path = os.path.join(work, f"{base}.{out}")
    preview = None
    if out == 'svg':
        svg = trace_to_svg(tr)
        with open(out_path, 'w', encoding='utf-8') as f:
            f.write(svg)
        preview = svg
    elif out == 'eps':
        with open(out_path, 'w', encoding='ascii') as f:
            f.write(trace_to_eps(tr))
        preview = trace_to_svg(tr)
    else:
        trace_to_dxf(tr, opts, out_path)
        preview = trace_to_svg(tr)
    notes = [f"{tr.count()} {'linhas' if any(l['kind'] == 'line' for l in tr.layers) else 'contornos'} "
             f"vetorizados em {len(tr.layers)} camada(s)."]
    if out == 'dxf':
        width_mm = _opt_float(opts, 'width_mm', 0.0, 0.0, 100000.0)
        k = (width_mm / tr.width) if width_mm > 0 else (25.4 / tr.dpi)
        notes.append(f"Tamanho no CAD: {tr.width * k:.1f} × {tr.height * k:.1f} mm.")
    return [out_path], {'preview_svg': preview, 'notes': notes}


# ---------------------------------------------------------------------------
# PDF, SVG, AI, EPS, XPS e eBooks (PyMuPDF / Ghostscript)
# ---------------------------------------------------------------------------

def _pymupdf():
    try:
        import pymupdf
    except ImportError:
        import fitz as pymupdf
    return pymupdf


def _gs_binary():
    return _which('gs', 'gswin64c', 'gswin32c')


def _ghostscript(args, what='O Ghostscript'):
    gs = _gs_binary()
    if not gs:
        raise ConversionError("Ghostscript não está instalado (sudo apt install ghostscript).")
    result = _run([gs, '-dSAFER', '-dBATCH', '-dNOPAUSE', '-dQUIET'] + args, timeout=300, what=what)
    if result.returncode != 0:
        raise ConversionError(f"{what} falhou: {_last_error_line(result, 'erro desconhecido')}")


def _as_pdf(src, in_ext, work):
    """Leva qualquer entrada "tipo página" para um PDF (caminho), mantendo os vetores."""
    pymupdf = _pymupdf()
    if in_ext == 'pdf':
        return src
    if in_ext in ('eps', 'ps'):
        out = os.path.join(work, '_intermediario.pdf')
        extra = ['-dEPSCrop'] if in_ext == 'eps' else []
        _ghostscript(['-sDEVICE=pdfwrite'] + extra + [f'-sOutputFile={out}', src])
        return out
    if in_ext == 'ai':
        try:
            with pymupdf.open(src, filetype='pdf') as d:
                if len(d):
                    return src
        except Exception:
            pass
        out = os.path.join(work, '_intermediario.pdf')
        _ghostscript(['-sDEVICE=pdfwrite', f'-sOutputFile={out}', src])
        return out
    try:
        with pymupdf.open(src, filetype=ALIASES.get(in_ext, in_ext)) as doc:
            pdf_bytes = doc.convert_to_pdf()
    except Exception as e:
        raise ConversionError(f"Não foi possível abrir o arquivo {in_ext.upper()} ({e}).")
    out = os.path.join(work, '_intermediario.pdf')
    with open(out, 'wb') as f:
        f.write(pdf_bytes)
    return out


def _open_pdf(pdf_path):
    pymupdf = _pymupdf()
    try:
        doc = pymupdf.open(pdf_path)
    except Exception as e:
        raise ConversionError(f"Não foi possível abrir o PDF ({e}).")
    if doc.needs_pass:
        raise ConversionError("O PDF está protegido por senha. Desbloqueie no app PDFs primeiro.")
    if not len(doc):
        raise ConversionError("O arquivo não tem nenhuma página.")
    return doc


def pages_to_raster(pdf_path, out, opts, work, base):
    from PIL import Image
    doc = _open_pdf(pdf_path)
    dpi = _opt_int(opts, 'dpi', 150, 36, 600)
    transparent = out in ('png', 'webp', 'tiff') and _opt_bool(opts, 'transparent')
    quality = _opt_int(opts, 'quality', 90, 1, 100)
    images = []
    for page in doc:
        pix = page.get_pixmap(dpi=dpi, alpha=transparent)
        mode = 'RGBA' if pix.alpha else 'RGB'
        images.append(Image.frombytes(mode, (pix.width, pix.height), pix.samples))
    doc.close()

    outputs = []
    if out == 'tiff' and len(images) > 1:
        path = os.path.join(work, f"{base}.tiff")
        images[0].save(path, save_all=True, append_images=images[1:], compression='tiff_lzw',
                       dpi=(dpi, dpi))
        return [path]
    for n, im in enumerate(images, 1):
        name = f"{base}.{out}" if len(images) == 1 else f"{base}_pagina_{n:03d}.{out}"
        path = os.path.join(work, name)
        kwargs = {}
        if out == 'jpg':
            im = _flatten(im)
            kwargs = {'quality': quality, 'optimize': True}
        elif out == 'webp':
            kwargs = {'quality': quality}
        elif out == 'gif':
            im = im.convert('RGB').quantize(256)
        if out in ('png', 'jpg', 'tiff', 'bmp'):
            kwargs['dpi'] = (dpi, dpi)
        im.save(path, format=PIL_FORMAT[out], **kwargs)
        outputs.append(path)
    return outputs


def pdf_to_svg(pdf_path, work, base):
    doc = _open_pdf(pdf_path)
    outputs = []
    for n, page in enumerate(doc, 1):
        svg = page.get_svg_image(text_as_path=True)
        name = f"{base}.svg" if len(doc) == 1 else f"{base}_pagina_{n:03d}.svg"
        path = os.path.join(work, name)
        with open(path, 'w', encoding='utf-8') as f:
            f.write(svg)
        outputs.append(path)
    doc.close()
    return outputs


def _is_white(c):
    return c is not None and all(v > 0.98 for v in c[:3])


def pdf_to_dxf(pdf_path, opts, work, base):
    """Extrai os vetores reais de cada página (linhas, Béziers, retângulos) e o
    texto, sem rasterizar - um PDF/SVG/AI exportado de CAD volta a ser geometria
    editável. Coordenadas em mm (1 pt = 25,4/72 mm)."""
    from ezdxf import path as ezpath
    from ezdxf.math import Vec2
    doc = _open_pdf(pdf_path)
    k = 25.4 / 72.0
    outputs = []
    notes = []
    total_vec = 0
    for pno, page in enumerate(doc, 1):
        dxf, version = _new_dxf(opts)
        msp = dxf.modelspace()
        for lname, aci in (('GEOMETRIA', 7), ('TEXTO', 7)):
            if lname not in dxf.layers:
                dxf.layers.add(lname, color=aci)
        x0, y0 = page.rect.x0, page.rect.y0
        H = page.rect.height

        def V(pt):
            return Vec2((pt.x - x0) * k, (H - (pt.y - y0)) * k)

        by_color = {}
        for d in page.get_drawings():
            stroke, fill = d.get('color'), d.get('fill')
            # Fundo branco sem contorno (comum em SVG/PDF) não é geometria
            if stroke is None and (fill is None or _is_white(fill)):
                continue
            color = stroke if stroke is not None else fill
            paths = []
            cur = None
            path = None
            for item in d['items']:
                op = item[0]
                if op in ('l', 'c'):
                    a = item[1]
                    if path is None or cur is None or abs(cur.x - a.x) > 1e-3 or abs(cur.y - a.y) > 1e-3:
                        path = ezpath.Path(V(a))
                        paths.append(path)
                    if op == 'l':
                        path.line_to(V(item[2]))
                        cur = item[2]
                    else:
                        path.curve4_to(V(item[4]), V(item[2]), V(item[3]))
                        cur = item[4]
                elif op == 're':
                    r = item[1]
                    p = ezpath.Path(V(r.tl))
                    for q in (r.tr, r.br, r.bl):
                        p.line_to(V(q))
                    p.close()
                    paths.append(p)
                    path, cur = None, None
                elif op == 'qu':
                    q = item[1]
                    p = ezpath.Path(V(q.ul))
                    for pt in (q.ur, q.lr, q.ll):
                        p.line_to(V(pt))
                    p.close()
                    paths.append(p)
                    path, cur = None, None
            if d.get('closePath') and path is not None and len(path):
                path.close()
            rgb = tuple(int(round(c * 255)) for c in (color or (0, 0, 0))[:3])
            by_color.setdefault(rgb, []).extend(p for p in paths if len(p))

        for rgb, paths in by_color.items():
            attribs = {'layer': 'GEOMETRIA'}
            if rgb not in ((0, 0, 0), (255, 255, 255)):
                attribs['true_color'] = _dxf_true_color(rgb)
            total_vec += len(paths)
            _render_paths(msp, paths, opts, attribs)

        if not _opt_bool(opts, 'skip_text'):
            text = page.get_text('dict')
            for block in text.get('blocks', []):
                for line in block.get('lines', []):
                    dx, dy = line.get('dir', (1, 0))
                    angle = math.degrees(math.atan2(-dy, dx))
                    for span in line.get('spans', []):
                        s = span.get('text', '').strip()
                        if not s:
                            continue
                        ox, oy = span['origin']
                        pos = Vec2((ox - x0) * k, (H - (oy - y0)) * k)
                        height = max(0.5, span.get('size', 10) * k * 0.72)
                        t = msp.add_text(s, height=height,
                                         dxfattribs={'layer': 'TEXTO', 'rotation': angle})
                        t.set_placement(pos)

        name = f"{base}.dxf" if len(doc) == 1 else f"{base}_pagina_{pno:03d}.dxf"
        path = os.path.join(work, name)
        _save_dxf(dxf, version, path)
        outputs.append(path)
    doc.close()
    if total_vec == 0:
        notes.append("Aviso: nenhum vetor foi encontrado - o arquivo parece conter só imagem "
                     "(ex.: PDF escaneado). Para vetorizar uma imagem, converta-a primeiro para "
                     "PNG e depois PNG → DXF.")
    else:
        notes.append(f"{total_vec} caminhos vetoriais extraídos (em mm).")
    return outputs, notes


def pdf_to_text(pdf_path, out, work, base):
    doc = _open_pdf(pdf_path)
    path = os.path.join(work, f"{base}.{out}")
    if out == 'txt':
        text = "\n\n".join(page.get_text() for page in doc)
        with open(path, 'w', encoding='utf-8') as f:
            f.write(text)
    else:
        body = "\n".join(page.get_text('xhtml') for page in doc)
        # O MuPDF escreve todo caractere não-ASCII como entidade (&#xe7;); volta
        # para UTF-8 legível, mantendo as entidades de marcação (&lt; &amp;...)
        body = re.sub(r'&#x([0-9a-fA-F]+);',
                      lambda m: chr(int(m.group(1), 16)) if int(m.group(1), 16) > 127 else m.group(0),
                      body)
        html_doc = ("<!doctype html><html lang=\"pt-br\"><head><meta charset=\"utf-8\">"
                    f"<title>{_html_escape(base)}</title><style>body{{max-width:820px;margin:40px auto;"
                    "font-family:Georgia,serif;line-height:1.6;padding:0 16px}img{max-width:100%}"
                    "</style></head><body>" + body + "</body></html>")
        with open(path, 'w', encoding='utf-8') as f:
            f.write(html_doc)
    doc.close()
    return [path]


def _html_escape(s):
    import html
    return html.escape(s or '')


# ---------------------------------------------------------------------------
# CAD (DXF / DWG)
# ---------------------------------------------------------------------------

def _read_dxf(path):
    import ezdxf
    from ezdxf import recover
    try:
        doc, auditor = recover.readfile(path)
    except IOError as e:
        raise ConversionError(f"Não foi possível ler o DXF ({e}).")
    except ezdxf.DXFStructureError as e:
        raise ConversionError(f"O DXF está corrompido de um jeito que não dá para recuperar ({e}).")
    return doc


def _dxf_page_render(doc, backend_cls, opts):
    from ezdxf.addons.drawing import Frontend, RenderContext, layout, config
    backend = backend_cls()
    bg = opts.get('cad_background', 'white')
    cfg = config.Configuration(
        background_policy=config.BackgroundPolicy.WHITE if bg == 'white' else config.BackgroundPolicy.BLACK,
        color_policy=config.ColorPolicy.COLOR,
        lineweight_scaling=1.0,
    )
    msp = doc.modelspace()
    Frontend(RenderContext(doc), backend, config=cfg).draw_layout(msp, finalize=True)
    page = layout.Page(0, 0, layout.Units.mm, margins=layout.Margins.all(5))
    return backend, page


def convert_dxf(src, out, opts, work, base):
    doc = _read_dxf(src)
    if len(doc.modelspace()) == 0:
        raise ConversionError("O desenho está vazio (nenhuma entidade no modelspace).")
    out_path = os.path.join(work, f"{base}.{out}")
    preview = None
    if out == 'dxf':
        version = DXF_VERSIONS.get(str(opts.get('dxf_version', 'R2010')).upper(), 'R2010')
        _finalize_dxf(doc)
        if version == 'R12':
            from ezdxf.addons import r12export
            r12export.saveas(doc, out_path)
        else:
            try:
                doc.dxfversion = {'R2000': 'AC1015', 'R2004': 'AC1018', 'R2007': 'AC1021',
                                  'R2010': 'AC1024', 'R2013': 'AC1027', 'R2018': 'AC1032'}[version]
            except Exception:
                pass
            doc.saveas(out_path)
        return [out_path], {'notes': [f"Salvo como DXF {version}."]}

    if out == 'svg':
        from ezdxf.addons.drawing import svg as svg_backend
        backend, page = _dxf_page_render(doc, svg_backend.SVGBackend, opts)
        svg = backend.get_string(page)
        with open(out_path, 'w', encoding='utf-8') as f:
            f.write(svg)
        preview = svg
        return [out_path], {'preview_svg': preview}

    from ezdxf.addons.drawing import pymupdf as pm_backend
    backend, page = _dxf_page_render(doc, pm_backend.PyMuPdfBackend, opts)
    if out in ('pdf', 'eps'):
        pdf_bytes = backend.get_pdf_bytes(page)
        pdf_path = out_path if out == 'pdf' else os.path.join(work, '_cad.pdf')
        with open(pdf_path, 'wb') as f:
            f.write(pdf_bytes)
        if out == 'eps':
            _ghostscript(['-sDEVICE=eps2write', f'-sOutputFile={out_path}', pdf_path])
        return [out_path], {}

    dpi = _opt_int(opts, 'dpi', 200, 36, 1200)
    png = backend.get_pixmap_bytes(page, fmt='png', dpi=dpi)
    if out == 'png':
        with open(out_path, 'wb') as f:
            f.write(png)
    else:
        from PIL import Image
        im = Image.open(io.BytesIO(png))
        im = _flatten(im) if out == 'jpg' else im
        kwargs = {'quality': _opt_int(opts, 'quality', 92, 1, 100)} if out in ('jpg', 'webp') else {}
        im.save(out_path, format=PIL_FORMAT[out], **kwargs)
    return [out_path], {}


ODA_VERSIONS = {'R2000': 'ACAD2000', 'R2004': 'ACAD2004', 'R2007': 'ACAD2007',
                'R2010': 'ACAD2010', 'R2013': 'ACAD2013', 'R2018': 'ACAD2018', 'R12': 'ACAD12'}


def _oda_convert(src, out_type, version, work):
    oda = _oda_converter()
    if not oda:
        raise ConversionError(unavailable_notes().get('dwg', 'ODA File Converter não encontrado.'))
    in_dir = os.path.join(work, '_oda_in')
    out_dir = os.path.join(work, '_oda_out')
    os.makedirs(in_dir, exist_ok=True)
    os.makedirs(out_dir, exist_ok=True)
    shutil.copy(src, in_dir)
    cmd = [oda, in_dir, out_dir, ODA_VERSIONS.get(version, 'ACAD2018'), out_type.upper(), '0', '1']
    env = dict(os.environ)
    env.setdefault('QT_QPA_PLATFORM', 'offscreen')
    try:
        subprocess.run(cmd, capture_output=True, timeout=300, env=env)
    except subprocess.TimeoutExpired as e:
        raise ConversionError("O ODA File Converter demorou demais.") from e
    for name in os.listdir(out_dir):
        if name.lower().endswith('.' + out_type.lower()):
            return os.path.join(out_dir, name)
    raise ConversionError("O ODA File Converter não gerou o arquivo de saída.")


def _dwg_to_dxf(src, work):
    if _oda_converter():
        return _oda_convert(src, 'dxf', 'R2018', work)
    tool = _libredwg_dwg2dxf()
    if not tool:
        raise ConversionError(unavailable_notes().get('dwg'))
    out = os.path.join(work, '_dwg.dxf')
    result = _run([tool, '-y', '-o', out, src], timeout=300, what='O LibreDWG')
    if not os.path.exists(out):
        raise ConversionError(f"O LibreDWG não conseguiu ler o DWG: {_last_error_line(result, '')}")
    return out


# ---------------------------------------------------------------------------
# LibreOffice (documentos, planilhas, apresentações, desenhos)
# ---------------------------------------------------------------------------

SOFFICE_TARGET = {
    'pdf': 'pdf', 'docx': 'docx:MS Word 2007 XML', 'doc': 'doc:MS Word 97', 'odt': 'odt',
    'rtf': 'rtf', 'txt': 'txt:Text (encoded):UTF8', 'html': 'html:XHTML Writer File:UTF8',
    'epub': 'epub', 'xlsx': 'xlsx:Calc MS Excel 2007 XML', 'xls': 'xls:MS Excel 97', 'ods': 'ods',
    'csv': 'csv:Text - txt - csv (StarCalc):44,34,76,1', 'pptx': 'pptx:Impress MS PowerPoint 2007 XML',
    'ppt': 'ppt:MS PowerPoint 97', 'odp': 'odp', 'svg': 'svg',
}


def _soffice(src, out, work, infilter=None, target=None):
    from Office import office_engine
    outdir = os.path.join(work, f'_lo_{out}')
    os.makedirs(outdir, exist_ok=True)
    try:
        result = office_engine.run_soffice(src, target or SOFFICE_TARGET.get(out, out), outdir,
                                           timeout=300, infilter=infilter)
    except office_engine.OfficeError as e:
        raise ConversionError(str(e))
    ext = out.split(':')[0]
    for name in os.listdir(outdir):
        if name.lower().endswith('.' + ext):
            return os.path.join(outdir, name)
    raise ConversionError(office_engine._soffice_error_message(result))


def _markdown_to_html(src, work, base):
    import markdown
    with open(src, 'r', encoding='utf-8', errors='replace') as f:
        text = f.read()
    body = markdown.markdown(text, extensions=['extra', 'sane_lists', 'toc'])
    html_doc = ("<!doctype html><html lang=\"pt-br\"><head><meta charset=\"utf-8\">"
                f"<title>{_html_escape(base)}</title><style>body{{max-width:820px;margin:40px auto;"
                "font-family:Arial,Helvetica,sans-serif;line-height:1.6;padding:0 16px}"
                "pre,code{background:#f4f4f4}pre{padding:12px;overflow:auto}"
                "table{border-collapse:collapse}td,th{border:1px solid #ccc;padding:4px 8px}"
                "</style></head><body>" + body + "</body></html>")
    path = os.path.join(work, f"{base}.html")
    with open(path, 'w', encoding='utf-8') as f:
        f.write(html_doc)
    return path


def _html_to_markdown(html_path, work, base):
    import html2text
    with open(html_path, 'r', encoding='utf-8', errors='replace') as f:
        html_text = f.read()
    h = html2text.HTML2Text()
    h.body_width = 0
    h.ignore_images = False
    md = h.handle(html_text)
    path = os.path.join(work, f"{base}.md")
    with open(path, 'w', encoding='utf-8') as f:
        f.write(md)
    return path


def convert_document(src, in_ext, out, opts, work, base):
    in_ext = ALIASES.get(in_ext, in_ext)
    notes = []
    # Markdown entra como HTML (o LibreOffice não lê Markdown)
    if in_ext == 'md':
        html_path = _markdown_to_html(src, work, base)
        if out == 'html':
            return [html_path], {}
        if out == 'txt':
            import html2text
            h = html2text.HTML2Text()
            h.ignore_links = True
            h.ignore_images = True
            h.ignore_emphasis = True
            h.body_width = 0
            path = os.path.join(work, f"{base}.txt")
            with open(html_path, encoding='utf-8') as f, open(path, 'w', encoding='utf-8') as g:
                g.write(h.handle(f.read()))
            return [path], {}
        src, in_ext = html_path, 'html'

    if out == 'md':
        html_path = src if in_ext == 'html' else _soffice(src, 'html', work)
        return [_html_to_markdown(html_path, work, base)], {}

    if out in ('png', 'jpg'):
        pdf = src if in_ext == 'pdf' else _soffice(src, 'pdf', work)
        return pages_to_raster(pdf, out, opts, work, base), {}

    infilter = None
    if in_ext == 'html' and out in ('pdf', 'docx', 'doc', 'odt', 'rtf', 'txt', 'epub'):
        infilter = 'HTML (StarWriter)'
    if in_ext == 'pdf':
        infilter = 'writer_pdf_import'
        notes.append("PDF → documento editável: o texto vira caixas posicionadas como no PDF. "
                     "Para só o texto corrido, use PDF → TXT.")
    produced = _soffice(src, out, work, infilter=infilter)
    final = os.path.join(work, f"{base}.{out}")
    shutil.move(produced, final)
    return [final], {'notes': notes}


def convert_draw(src, in_ext, out, opts, work, base):
    """CorelDRAW, Visio, WMF/EMF, ODG: o LibreOffice Draw gera um PDF vetorial e
    o resto segue o caminho do PDF (inclusive PDF -> DXF com geometria real)."""
    if out == 'svg':
        produced = _soffice(src, 'svg', work, target='svg:draw_svg_Export')
        final = os.path.join(work, f"{base}.svg")
        shutil.move(produced, final)
        with open(final, encoding='utf-8', errors='ignore') as f:
            preview = f.read()
        return [final], {'preview_svg': preview if len(preview) < 3_000_000 else None}
    pdf = _soffice(src, 'pdf', work, target='pdf:draw_pdf_Export')
    return convert_pdflike(pdf, 'pdf', out, opts, work, base)


# ---------------------------------------------------------------------------
# Dados e planilhas (CSV/TSV/JSON/YAML/XML/XLSX)
# ---------------------------------------------------------------------------

def _read_text(path):
    with open(path, 'rb') as f:
        raw = f.read()
    for enc in ('utf-8-sig', 'cp1252', 'latin-1'):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode('utf-8', 'replace')


def _xml_to_obj(elem):
    children = list(elem)
    node = {f"@{k}": v for k, v in elem.attrib.items()}
    if children:
        for child in children:
            tag = child.tag.split('}')[-1]
            value = _xml_to_obj(child)
            if tag in node:
                if not isinstance(node[tag], list):
                    node[tag] = [node[tag]]
                node[tag].append(value)
            else:
                node[tag] = value
        text = (elem.text or '').strip()
        if text:
            node['#text'] = text
        return node
    text = (elem.text or '').strip()
    if node:
        if text:
            node['#text'] = text
        return node
    return text


def _obj_to_xml(parent, key, value):
    from lxml import etree
    tag = re.sub(r'[^\w.-]', '_', str(key)) or 'item'
    if tag[0].isdigit() or tag[0] in '.-':
        tag = 'item_' + tag
    if isinstance(value, list):
        for v in value:
            _obj_to_xml(parent, key, v)
        return
    el = etree.SubElement(parent, tag)
    if isinstance(value, dict):
        for k, v in value.items():
            if str(k).startswith('@'):
                el.set(re.sub(r'[^\w.-]', '_', str(k)[1:]) or 'attr', '' if v is None else str(v))
            elif k == '#text':
                el.text = '' if v is None else str(v)
            else:
                _obj_to_xml(el, k, v)
    else:
        el.text = '' if value is None else str(value)


def _flatten_record(obj, prefix=''):
    out = {}
    if isinstance(obj, dict):
        for k, v in obj.items():
            key = f"{prefix}.{k}" if prefix else str(k)
            if isinstance(v, (dict, list)):
                out.update(_flatten_record(v, key))
            else:
                out[key] = v
    elif isinstance(obj, list):
        if all(not isinstance(v, (dict, list)) for v in obj):
            out[prefix or 'valor'] = ", ".join('' if v is None else str(v) for v in obj)
        else:
            for i, v in enumerate(obj):
                out.update(_flatten_record(v, f"{prefix}[{i}]" if prefix else f"[{i}]"))
    else:
        out[prefix or 'valor'] = obj
    return out


def _tree_to_tables(obj):
    """Transforma JSON/YAML/XML em tabelas {nome: [cabeçalho, linhas...]}."""
    records = None
    if isinstance(obj, list):
        records = obj
    elif isinstance(obj, dict):
        # {"clientes": [...]} ou XML <root><row>..</row><row>..</row></root>
        lists = [(k, v) for k, v in obj.items() if isinstance(v, list)]
        if len(obj) == 1:
            only = next(iter(obj.values()))
            if isinstance(only, dict):
                inner = [(k, v) for k, v in only.items() if isinstance(v, list)]
                if len(inner) == 1 and len(only) == 1:
                    records = inner[0][1]
                else:
                    records = [only]
            elif isinstance(only, list):
                records = only
            else:
                records = [obj]
        elif len(lists) >= 1 and all(isinstance(v, list) for v in obj.values()):
            tables = {}
            for k, v in lists:
                tables.update({str(k)[:31]: t for t in _tree_to_tables(v).values()})
            return tables
        else:
            records = [obj]
    else:
        records = [{'valor': obj}]
    flat = [_flatten_record(r) if isinstance(r, (dict, list)) else {'valor': r} for r in records]
    header = []
    for r in flat:
        for k in r:
            if k not in header:
                header.append(k)
    rows = [[r.get(h, '') for h in header] for r in flat]
    return {'Planilha1': [header] + rows}


def _tables_to_tree(tables):
    def to_records(rows):
        if not rows:
            return []
        header = [str(h) if h not in (None, '') else f"coluna_{i + 1}" for i, h in enumerate(rows[0])]
        # CSV só tem texto: "30" vira 30 (sem estragar CEP/telefone com zero à esquerda)
        return [dict(zip(header, [_json_safe(_cell(c) if isinstance(c, str) else c) for c in row]))
                for row in rows[1:]]
    if len(tables) == 1:
        return to_records(next(iter(tables.values())))
    return {name: to_records(rows) for name, rows in tables.items()}


def _json_safe(v):
    import datetime
    if isinstance(v, (datetime.date, datetime.datetime, datetime.time)):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def _load_data(src, in_ext, work):
    """Devolve ('tree', obj) ou ('tables', {nome: linhas})."""
    in_ext = ALIASES.get(in_ext, in_ext)
    if in_ext == 'json':
        text = _read_text(src)
        try:
            return 'tree', json.loads(text)
        except json.JSONDecodeError:
            # JSON Lines (um objeto por linha)
            try:
                return 'tree', [json.loads(l) for l in text.splitlines() if l.strip()]
            except json.JSONDecodeError as e:
                raise ConversionError(f"JSON inválido: {e}")
    if in_ext == 'yaml':
        import yaml
        try:
            docs = [d for d in yaml.safe_load_all(_read_text(src))]
        except yaml.YAMLError as e:
            raise ConversionError(f"YAML inválido: {e}")
        return 'tree', docs[0] if len(docs) == 1 else docs
    if in_ext == 'xml':
        from lxml import etree
        try:
            root = etree.parse(src, etree.XMLParser(resolve_entities=False, no_network=True,
                                                    huge_tree=True)).getroot()
        except etree.XMLSyntaxError as e:
            raise ConversionError(f"XML inválido: {e}")
        return 'tree', {root.tag.split('}')[-1]: _xml_to_obj(root)}
    if in_ext in ('csv', 'tsv'):
        text = _read_text(src)
        if in_ext == 'tsv':
            delim = '\t'
        else:
            try:
                delim = csv.Sniffer().sniff(text[:20000], delimiters=',;\t|').delimiter
            except csv.Error:
                delim = ','
        rows = list(csv.reader(io.StringIO(text), delimiter=delim))
        return 'tables', {'Planilha1': rows}
    if in_ext in ('xls', 'ods', 'numbers'):
        src = _soffice(src, 'xlsx', work)
        in_ext = 'xlsx'
    if in_ext == 'xlsx':
        import openpyxl
        try:
            wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
        except Exception as e:
            raise ConversionError(f"Não foi possível ler a planilha ({e}).")
        tables = {}
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            while rows and all(c in (None, '') for c in rows[-1]):
                rows.pop()
            width = max((max((i + 1 for i, c in enumerate(r) if c not in (None, '')), default=0)
                         for r in rows), default=0)
            tables[ws.title] = [[('' if c is None else c) for c in r[:width]] for r in rows]
        wb.close()
        return 'tables', tables
    raise ConversionError(f"Leitura de {in_ext.upper()} como dados não é suportada.")


def _write_xlsx(tables, path):
    import openpyxl
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in tables.items():
        safe = re.sub(r'[\[\]:*?/\\]', '_', str(name))[:31] or 'Planilha'
        ws = wb.create_sheet(safe)
        for row in rows:
            ws.append([_cell(v) for v in row])
        for i, col in enumerate(ws.iter_cols(max_row=min(ws.max_row, 200)), 1):
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(8, width + 2))
        if rows:
            from openpyxl.styles import Font
            for c in ws[1]:
                c.font = Font(bold=True)
    if not wb.worksheets:
        wb.create_sheet('Planilha1')
    wb.save(path)


def _cell(v):
    if isinstance(v, str):
        s = v.strip()
        if re.fullmatch(r'-?\d{1,15}', s) and not (len(s) > 1 and s.lstrip('-').startswith('0')):
            return int(s)
        if re.fullmatch(r'-?\d+\.\d+', s):
            try:
                return float(s)
            except ValueError:
                return v
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)
    return v


def convert_data(src, in_ext, out, opts, work, base):
    kind, payload = _load_data(src, in_ext, work)
    path = os.path.join(work, f"{base}.{out}")
    tables = payload if kind == 'tables' else None
    tree = payload if kind == 'tree' else None

    if out in ('json', 'yaml', 'xml'):
        obj = tree if tree is not None else _tables_to_tree(tables)
        if out == 'json':
            with open(path, 'w', encoding='utf-8') as f:
                json.dump(obj, f, ensure_ascii=False, indent=2, default=str)
        elif out == 'yaml':
            import yaml
            with open(path, 'w', encoding='utf-8') as f:
                yaml.safe_dump(json.loads(json.dumps(obj, default=str)), f, allow_unicode=True,
                               sort_keys=False)
        else:
            from lxml import etree
            if isinstance(obj, dict) and len(obj) == 1 and not isinstance(next(iter(obj.values())), list):
                key, value = next(iter(obj.items()))
                holder = etree.Element('tmp')
                _obj_to_xml(holder, key, value)
                root = holder[0]
            else:
                root = etree.Element('dados')
                if isinstance(obj, list):
                    for item in obj:
                        _obj_to_xml(root, 'registro', item)
                else:
                    for k, v in obj.items():
                        _obj_to_xml(root, k, v)
            etree.ElementTree(root).write(path, encoding='utf-8', xml_declaration=True,
                                          pretty_print=True)
        return [path], {}

    if tables is None:
        tables = _tree_to_tables(tree)

    if out in ('csv', 'tsv'):
        outputs = []
        for n, (name, rows) in enumerate(tables.items()):
            p = path if len(tables) == 1 else os.path.join(work, f"{base}_{_safe_base(name)}.{out}")
            with open(p, 'w', encoding='utf-8-sig' if out == 'csv' else 'utf-8', newline='') as f:
                writer = csv.writer(f, delimiter=',' if out == 'csv' else '\t')
                for row in rows:
                    writer.writerow(['' if c is None else (json.dumps(c, ensure_ascii=False)
                                                           if isinstance(c, (dict, list)) else c)
                                     for c in row])
            outputs.append(p)
        return outputs, {}

    if out == 'html':
        parts = ["<!doctype html><html lang=\"pt-br\"><head><meta charset=\"utf-8\">",
                 f"<title>{_html_escape(base)}</title><style>body{{font-family:Arial,sans-serif;margin:24px}}"
                 "table{border-collapse:collapse;margin-bottom:32px}th,td{border:1px solid #ccc;"
                 "padding:4px 10px;text-align:left}th{background:#f0f0f0}</style></head><body>"]
        for name, rows in tables.items():
            if len(tables) > 1:
                parts.append(f"<h2>{_html_escape(str(name))}</h2>")
            parts.append("<table>")
            for i, row in enumerate(rows):
                tag = 'th' if i == 0 else 'td'
                parts.append("<tr>" + "".join(f"<{tag}>{_html_escape('' if c is None else str(c))}</{tag}>"
                                              for c in row) + "</tr>")
            parts.append("</table>")
        parts.append("</body></html>")
        with open(path, 'w', encoding='utf-8') as f:
            f.write("".join(parts))
        return [path], {}

    xlsx_path = os.path.join(work, f"{base}.xlsx")
    _write_xlsx(tables, xlsx_path)
    if out == 'xlsx':
        return [xlsx_path], {}
    produced = _soffice(xlsx_path, out, work)
    final = os.path.join(work, f"{base}_final.{out}")
    shutil.move(produced, final)
    os.remove(xlsx_path)
    target = os.path.join(work, f"{base}.{out}")
    shutil.move(final, target)
    return [target], {}


def convert_spreadsheet(src, in_ext, out, opts, work, base):
    in_ext = ALIASES.get(in_ext, in_ext)
    # Formatos de planilha "de verdade" passam pelo LibreOffice para manter
    # fórmulas, formatação e várias abas
    if out in ('xlsx', 'xls', 'ods', 'pdf') and in_ext in ('xlsx', 'xls', 'ods', 'numbers'):
        produced = _soffice(src, out, work)
        final = os.path.join(work, f"{base}.{out}")
        shutil.move(produced, final)
        return [final], {}
    if out in ('pdf', 'ods', 'xls') and in_ext == 'csv':
        produced = _soffice(src, out, work, infilter='Text - txt - csv (StarCalc)')
        final = os.path.join(work, f"{base}.{out}")
        shutil.move(produced, final)
        return [final], {}
    return convert_data(src, in_ext, out, opts, work, base)


def convert_slides(src, in_ext, out, opts, work, base):
    if out in ('png', 'jpg'):
        pdf = _soffice(src, 'pdf', work)
        return pages_to_raster(pdf, out, opts, work, base), {}
    produced = _soffice(src, out, work)
    final = os.path.join(work, f"{base}.{out}")
    shutil.move(produced, final)
    return [final], {}


# ---------------------------------------------------------------------------
# Áudio e vídeo (FFmpeg)
# ---------------------------------------------------------------------------

def _audio_args(out, opts):
    br = _opt_int(opts, 'audio_bitrate', 192, 32, 320)
    return {
        'mp3': ['-c:a', 'libmp3lame', '-b:a', f'{br}k'],
        'wav': ['-c:a', 'pcm_s16le'],
        'flac': ['-c:a', 'flac'],
        'aac': ['-c:a', 'aac', '-b:a', f'{br}k'],
        'm4a': ['-c:a', 'aac', '-b:a', f'{br}k'],
        'ogg': ['-c:a', 'libvorbis', '-b:a', f'{br}k'],
        'opus': ['-c:a', 'libopus', '-b:a', f'{min(br, 256)}k'],
        'wma': ['-c:a', 'wmav2', '-b:a', f'{br}k'],
        'aiff': ['-c:a', 'pcm_s16be'],
        'ac3': ['-c:a', 'ac3', '-b:a', f'{max(br, 192)}k'],
        'amr': ['-c:a', 'libopencore_amrnb', '-ar', '8000', '-ac', '1', '-b:a', '12.2k'],
        'mp2': ['-c:a', 'mp2', '-b:a', f'{max(br, 128)}k'],
    }[out]


def _video_args(out, opts):
    crf = {'alta': '19', 'media': '23', 'baixa': '28'}.get(opts.get('video_quality', 'media'), '23')
    x264 = ['-c:v', 'libx264', '-preset', 'medium', '-crf', crf, '-pix_fmt', 'yuv420p']
    return {
        'mp4': x264 + ['-c:a', 'aac', '-b:a', '192k', '-movflags', '+faststart'],
        'm4v': x264 + ['-c:a', 'aac', '-b:a', '192k', '-movflags', '+faststart'],
        'mov': x264 + ['-c:a', 'aac', '-b:a', '192k', '-movflags', '+faststart'],
        'mkv': x264 + ['-c:a', 'aac', '-b:a', '192k'],
        'ts': x264 + ['-c:a', 'aac', '-b:a', '192k'],
        'flv': x264 + ['-c:a', 'aac', '-b:a', '128k', '-ar', '44100'],
        '3gp': ['-c:v', 'libx264', '-profile:v', 'baseline', '-level', '3.0', '-crf', crf,
                '-pix_fmt', 'yuv420p', '-c:a', 'aac', '-b:a', '96k', '-ac', '1', '-ar', '22050'],
        'webm': ['-c:v', 'libvpx-vp9', '-crf', str(int(crf) + 9), '-b:v', '0', '-row-mt', '1',
                 '-deadline', 'good', '-cpu-used', '4', '-pix_fmt', 'yuv420p', '-c:a', 'libopus', '-b:a', '128k'],
        'avi': ['-c:v', 'libxvid', '-q:v', '4', '-c:a', 'libmp3lame', '-b:a', '192k'],
        'wmv': ['-c:v', 'wmv2', '-q:v', '4', '-c:a', 'wmav2', '-b:a', '192k'],
        'mpg': ['-c:v', 'mpeg2video', '-q:v', '3', '-c:a', 'mp2', '-b:a', '224k'],
        'ogv': ['-c:v', 'libtheora', '-q:v', '7', '-c:a', 'libvorbis', '-q:a', '5'],
    }[out]


def convert_media(src, in_ext, out, opts, work, base):
    ffmpeg = _which('ffmpeg')
    if not ffmpeg:
        raise ConversionError(unavailable_notes().get('mp4'))
    out_path = os.path.join(work, f"{base}.{out}")
    in_group = FORMATS.get(in_ext, ('',))[0]
    cmd = [ffmpeg, '-hide_banner', '-y']
    height = _opt_int(opts, 'video_height', 0, 0, 4320)

    if in_group == 'audio' and out in VIDEO_OUT:
        # Áudio -> vídeo com tela preta (ex.: subir um áudio no YouTube)
        cmd += ['-f', 'lavfi', '-i', 'color=c=black:s=1280x720:r=2', '-i', src, '-shortest',
                '-c:v', 'libx264', '-tune', 'stillimage', '-pix_fmt', 'yuv420p',
                '-c:a', 'aac', '-b:a', f"{_opt_int(opts, 'audio_bitrate', 192, 32, 320)}k", out_path]
    elif out == 'gif':
        fps = _opt_int(opts, 'gif_fps', 12, 1, 30)
        width = _opt_int(opts, 'gif_width', 480, 64, 1920)
        vf = (f"fps={fps},scale='min({width},iw)':-2:flags=lanczos,split[s0][s1];"
              f"[s0]palettegen=stats_mode=diff[p];[s1][p]paletteuse=dither=bayer:bayer_scale=4")
        cmd += ['-i', src, '-vf', vf, '-loop', '0', out_path]
    elif out in AUDIO_OUT:
        cmd += ['-i', src, '-vn', '-map', '0:a:0?'] + _audio_args(out, opts) + [out_path]
    else:
        vf = f"scale=-2:'min({height},ih)'" if height else "scale=trunc(iw/2)*2:trunc(ih/2)*2"
        if in_ext == 'gif':
            # GIF não tem áudio: só a parte de vídeo dos argumentos
            args = _video_args(out, opts)
            if '-c:a' in args:
                args = args[:args.index('-c:a')]
            cmd += ['-i', src, '-vf', vf] + args + ['-an', out_path]
        else:
            cmd += ['-i', src, '-map', '0:v:0', '-map', '0:a?', '-vf', vf] + _video_args(out, opts) + [out_path]

    result = _run(cmd, timeout=3600, what='O FFmpeg')
    if result.returncode != 0 or not os.path.exists(out_path) or os.path.getsize(out_path) == 0:
        detail = _last_error_line(result, 'erro desconhecido do FFmpeg')
        if 'does not contain any stream' in detail or 'Output file is empty' in detail \
                or 'matches no streams' in detail:
            detail = 'o arquivo não tem faixa de áudio/vídeo compatível com esta saída'
        raise ConversionError(f"O FFmpeg não conseguiu converter para {out.upper()}: {detail}")
    return [out_path], {}


# ---------------------------------------------------------------------------
# Compactados (zip/tar nativos, 7-Zip para o resto)
# ---------------------------------------------------------------------------

def _extract_archive(src, in_ext, dest):
    in_ext = ALIASES.get(in_ext, in_ext)
    os.makedirs(dest, exist_ok=True)

    def safe_members_ok(names):
        for n in names:
            norm = os.path.normpath(n)
            if norm.startswith('..') or os.path.isabs(norm):
                raise ConversionError("O arquivo compactado tem caminhos inseguros (../) e foi recusado.")

    if in_ext in ('zip', 'jar', 'apk') and zipfile.is_zipfile(src):
        with zipfile.ZipFile(src) as zf:
            if any(i.flag_bits & 0x1 for i in zf.infolist()):
                raise ConversionError("O ZIP é protegido por senha.")
            safe_members_ok(zf.namelist())
            zf.extractall(dest)
        return
    if in_ext in ('tar', 'tar.gz', 'tar.bz2', 'tar.xz'):
        try:
            with tarfile.open(src) as tf:
                members = [m for m in tf.getmembers() if m.isfile() or m.isdir()]
                safe_members_ok([m.name for m in members])
                tf.extractall(dest, members=members, filter='data')
            return
        except tarfile.TarError as e:
            raise ConversionError(f"TAR inválido: {e}")
    if in_ext in ('gz', 'bz2', 'xz'):
        import gzip, bz2, lzma
        opener = {'gz': gzip.open, 'bz2': bz2.open, 'xz': lzma.open}[in_ext]
        inner = os.path.splitext(os.path.basename(src))[0] or 'arquivo'
        try:
            with opener(src, 'rb') as fin, open(os.path.join(dest, inner), 'wb') as fout:
                shutil.copyfileobj(fin, fout)
        except (OSError, EOFError, lzma.LZMAError) as e:
            raise ConversionError(f"Não foi possível descompactar ({e}).")
        # .gz que na verdade é .tar.gz
        inner_path = os.path.join(dest, inner)
        if tarfile.is_tarfile(inner_path):
            sub = os.path.join(dest, '_tar')
            with tarfile.open(inner_path) as tf:
                tf.extractall(sub, filter='data')
            os.remove(inner_path)
            for n in os.listdir(sub):
                shutil.move(os.path.join(sub, n), os.path.join(dest, n))
            os.rmdir(sub)
        return
    seven = _which('7z', '7za', '7zz')
    if not seven:
        raise ConversionError("Para abrir este formato instale o 7-Zip (sudo apt install 7zip).")
    result = _run([seven, 'x', '-y', '-p', f'-o{dest}', src], timeout=900, what='O 7-Zip')
    if result.returncode != 0:
        detail = _last_error_line(result, 'erro desconhecido')
        if 'Wrong password' in detail or 'password' in detail.lower():
            raise ConversionError("O arquivo compactado é protegido por senha.")
        raise ConversionError(f"O 7-Zip não conseguiu abrir o arquivo: {detail}")
    # .deb/.rpm e afins trazem um .tar dentro: abre mais uma camada
    items = os.listdir(dest)
    if len(items) == 1 and items[0].endswith('.tar'):
        inner = os.path.join(dest, items[0])
        with tarfile.open(inner) as tf:
            tf.extractall(dest, filter='data')
        os.remove(inner)


def convert_archive(src, in_ext, out, opts, work, base):
    tree = os.path.join(work, '_extraido')
    _extract_archive(src, in_ext, tree)
    files = []
    for root, _, names in os.walk(tree):
        for n in names:
            full = os.path.join(root, n)
            files.append((full, os.path.relpath(full, tree)))
    if not files:
        raise ConversionError("O arquivo compactado está vazio.")
    out_path = os.path.join(work, f"{base}.{out}")
    if out == 'zip':
        with zipfile.ZipFile(out_path, 'w', zipfile.ZIP_DEFLATED, compresslevel=9) as zf:
            for full, rel in files:
                zf.write(full, rel)
    elif out == '7z':
        seven = _which('7z', '7za', '7zz')
        try:
            result = subprocess.run([seven, 'a', '-t7z', '-mx=7', out_path, '.'],
                                    cwd=tree, capture_output=True, timeout=1800)
        except subprocess.TimeoutExpired as e:
            raise ConversionError("O 7-Zip demorou demais (tempo limite excedido).") from e
        if result.returncode != 0:
            raise ConversionError(f"O 7-Zip falhou: {_last_error_line(result, '')}")
    else:
        mode = {'tar': 'w', 'tar.gz': 'w:gz', 'tar.bz2': 'w:bz2', 'tar.xz': 'w:xz'}[out]
        with tarfile.open(out_path, mode) as tf:
            for full, rel in files:
                tf.add(full, arcname=rel)
    return [out_path], {'notes': [f"{len(files)} arquivo(s) reempacotado(s)."]}


# ---------------------------------------------------------------------------
# Fontes (fontTools)
# ---------------------------------------------------------------------------

def _otf_to_ttf(font):
    """CFF (cúbico) -> glyf (quadrático). Receita oficial do fontTools (otf2ttf)."""
    from fontTools.pens.cu2quPen import Cu2QuPen
    from fontTools.pens.ttGlyphPen import TTGlyphPen
    from fontTools.ttLib import newTable

    glyph_order = font.getGlyphOrder()
    glyph_set = font.getGlyphSet()
    quad = {}
    for name in glyph_order:
        tt_pen = TTGlyphPen(glyph_set)
        glyph_set[name].draw(Cu2QuPen(tt_pen, 1.0, reverse_direction=True))
        quad[name] = tt_pen.glyph()

    font['loca'] = newTable('loca')
    font['glyf'] = glyf = newTable('glyf')
    glyf.glyphOrder = glyph_order
    glyf.glyphs = quad
    for tag in ('CFF ', 'CFF2', 'VORG'):
        if tag in font:
            del font[tag]
    glyf.compile(font)
    hmtx = font['hmtx']
    for name, g in glyf.glyphs.items():
        if hasattr(g, 'xMin'):
            hmtx[name] = (hmtx[name][0], g.xMin)
    font['maxp'] = maxp = newTable('maxp')
    maxp.tableVersion = 0x00010000
    for attr in ('maxZones',):
        setattr(maxp, attr, 1)
    for attr in ('maxTwilightPoints', 'maxStorage', 'maxFunctionDefs', 'maxInstructionDefs',
                 'maxStackElements', 'maxSizeOfInstructions'):
        setattr(maxp, attr, 0)
    maxp.maxComponentElements = max(
        (len(g.components) if hasattr(g, 'components') else 0) for g in glyf.glyphs.values())
    maxp.compile(font)
    post = font['post']
    post.formatType = 2.0
    post.extraNames = []
    post.mapping = {}
    post.glyphOrder = glyph_order
    try:
        post.compile(font)
    except OverflowError:
        post.formatType = 3
    font.sfntVersion = '\x00\x01\x00\x00'


def _ttf_to_otf(font):
    """glyf (quadrático) -> CFF (cúbico) desenhando cada glifo com T2CharStringPen."""
    from fontTools.fontBuilder import FontBuilder
    from fontTools.pens.t2CharStringPen import T2CharStringPen

    glyph_order = font.getGlyphOrder()
    glyph_set = font.getGlyphSet()
    hmtx = font['hmtx']
    charstrings = {}
    for name in glyph_order:
        pen = T2CharStringPen(hmtx[name][0], glyph_set)
        glyph_set[name].draw(pen)
        charstrings[name] = pen.getCharString()

    name_table = font['name']
    ps_name = (name_table.getDebugName(6) or 'Font').replace(' ', '')
    font_info = {
        'FullName': name_table.getDebugName(4) or ps_name,
        'FamilyName': name_table.getDebugName(1) or ps_name,
        'Weight': 'Regular',
    }
    for tag in ('glyf', 'loca', 'fpgm', 'prep', 'cvt ', 'gasp', 'hdmx', 'LTSH', 'VDMX'):
        if tag in font:
            del font[tag]
    fb = FontBuilder(font=font, isTTF=False)
    fb.setupCFF(ps_name, font_info, charstrings, {})
    font['maxp'].tableVersion = 0x00005000
    font.sfntVersion = 'OTTO'
    # No CFF o lsb de cada glifo tem que bater com o xMin real do contorno
    for name in glyph_order:
        bounds = charstrings[name].calcBounds(None)
        hmtx[name] = (hmtx[name][0], int(round(bounds[0])) if bounds else 0)


def convert_font(src, in_ext, out, opts, work, base):
    from fontTools.ttLib import TTFont
    try:
        font = TTFont(src)
    except Exception as e:
        raise ConversionError(f"Não foi possível ler a fonte ({e}).")
    notes = []
    is_cff = 'CFF ' in font or 'CFF2' in font
    try:
        if out == 'ttf' and is_cff:
            if 'CFF2' in font:
                raise ConversionError("Fontes variáveis CFF2 não podem ser convertidas para TTF.")
            _otf_to_ttf(font)
            notes.append("Contornos CFF convertidos para TrueType (quadráticos).")
        elif out == 'otf' and not is_cff:
            if 'fvar' in font:
                raise ConversionError("Fontes variáveis TrueType não podem virar OTF (CFF). "
                                      "Converta para WOFF2 ou TTF.")
            _ttf_to_otf(font)
            notes.append("Contornos TrueType convertidos para CFF (OpenType PostScript).")
        font.flavor = {'woff': 'woff', 'woff2': 'woff2'}.get(out)
        out_path = os.path.join(work, f"{base}.{out}")
        font.save(out_path)
    except ConversionError:
        raise
    except Exception as e:
        raise ConversionError(f"O fontTools não conseguiu gerar {out.upper()}: {e}")
    return [out_path], {'notes': notes}


# ---------------------------------------------------------------------------
# 3D (trimesh)
# ---------------------------------------------------------------------------

def convert_model(src, in_ext, out, opts, work, base):
    import trimesh
    try:
        mesh = trimesh.load(src, file_type=in_ext, force='mesh', process=False)
    except Exception as e:
        raise ConversionError(f"Não foi possível ler o modelo 3D ({e}).")
    if mesh is None or getattr(mesh, 'faces', None) is None or len(mesh.faces) == 0:
        raise ConversionError("O modelo 3D não tem nenhuma face.")
    out_path = os.path.join(work, f"{base}.{out}")
    notes = [f"{len(mesh.faces)} faces, {len(mesh.vertices)} vértices."]
    if out == 'dxf':
        doc, version = _new_dxf(opts)
        msp = doc.modelspace()
        doc.layers.add('MALHA_3D', color=7)
        tris = mesh.triangles
        if len(tris) > 400000:
            raise ConversionError("Modelo grande demais para DXF (mais de 400 mil faces).")
        for t in tris:
            a, b, c = (tuple(map(float, p)) for p in t)
            msp.add_3dface([a, b, c, c], dxfattribs={'layer': 'MALHA_3D'})
        _save_dxf(doc, version, out_path)
        notes.append("Cada triângulo virou uma entidade 3DFACE (unidades do modelo).")
        return [out_path], {'notes': notes}
    try:
        data = mesh.export(file_type=out)
    except Exception as e:
        raise ConversionError(f"O trimesh não conseguiu gerar {out.upper()}: {e}")
    mode = 'w' if isinstance(data, str) else 'wb'
    with open(out_path, mode, **({'encoding': 'utf-8'} if mode == 'w' else {})) as f:
        f.write(data)
    return [out_path], {'notes': notes}


# ---------------------------------------------------------------------------
# PDF-like: roteamento das saídas
# ---------------------------------------------------------------------------

def convert_pdflike(src, in_ext, out, opts, work, base):
    in_ext = ALIASES.get(in_ext, in_ext)
    extras = {}
    pdf = _as_pdf(src, in_ext, work)
    if out in PAGE_RASTER_OUT:
        return pages_to_raster(pdf, out, opts, work, base), extras
    if out == 'svg':
        outs = pdf_to_svg(pdf, work, base)
        if len(outs) == 1:
            with open(outs[0], encoding='utf-8') as f:
                svg = f.read()
            if len(svg) < 3_000_000:
                extras['preview_svg'] = svg
        return outs, extras
    if out == 'dxf':
        outs, notes = pdf_to_dxf(pdf, opts, work, base)
        extras['notes'] = notes
        if len(outs) == 1:
            try:
                extras['preview_svg'] = _dxf_preview(outs[0], opts)
            except Exception:
                pass
        return outs, extras
    if out == 'pdf':
        final = os.path.join(work, f"{base}.pdf")
        shutil.copy(pdf, final)
        return [final], extras
    if out in ('eps', 'ps'):
        final = os.path.join(work, f"{base}.{out}")
        doc = _open_pdf(pdf)
        n = len(doc)
        doc.close()
        args = ['-sDEVICE=eps2write' if out == 'eps' else '-sDEVICE=ps2write']
        if out == 'eps' and n > 1:
            args += ['-dFirstPage=1', '-dLastPage=1']
            extras['notes'] = ["EPS guarda uma página só: foi usada a primeira."]
        _ghostscript(args + [f'-sOutputFile={final}', pdf])
        return [final], extras
    if out in ('txt', 'html'):
        return pdf_to_text(pdf, out, work, base), extras
    if out in ('docx', 'odt'):
        return convert_document(pdf, 'pdf', out, opts, work, base)
    raise ConversionError(f"Conversão de {in_ext.upper()} para {out.upper()} não suportada.")


def _dxf_preview(dxf_path, opts):
    from ezdxf.addons.drawing import svg as svg_backend
    doc = _read_dxf(dxf_path)
    if len(doc.modelspace()) > 30000:
        return None
    backend, page = _dxf_page_render(doc, svg_backend.SVGBackend, opts)
    svg = backend.get_string(page)
    return svg if len(svg) < 3_000_000 else None


def convert_ebook(src, in_ext, out, opts, work, base):
    in_ext = ALIASES.get(in_ext, in_ext)
    if out in ('docx', 'odt', 'epub'):
        pdf = _as_pdf(src, in_ext, work)
        html_path = pdf_to_text(pdf, 'html', work, '_livro')[0]
        produced = _soffice(html_path, out, work, infilter='HTML (StarWriter)')
        final = os.path.join(work, f"{base}.{out}")
        shutil.move(produced, final)
        return [final], {}
    return convert_pdflike(src, in_ext, out, opts, work, base)


# ---------------------------------------------------------------------------
# Ponto de entrada
# ---------------------------------------------------------------------------

def _dispatch(src, in_ext, out, opts, work, base):
    group = FORMATS.get(in_ext, ('',))[0]
    real = ALIASES.get(in_ext, in_ext)

    if group == 'image':
        if out in TRACE_OUT:
            return convert_trace(src, out, opts, work, base)
        if out in VIDEO_OUT:
            return convert_media(src, in_ext, out, opts, work, base)
        return convert_raster(src, out, opts, work, base), {}
    if real in ('pdf', 'ai', 'eps', 'ps', 'svg', 'xps'):
        return convert_pdflike(src, in_ext, out, opts, work, base)
    if real in EBOOK_IN:
        return convert_ebook(src, in_ext, out, opts, work, base)
    if real == 'dxf':
        if out == 'dwg':
            version = str(opts.get('dwg_version', 'R2018')).upper()
            produced = _oda_convert(src, 'dwg', version, work)
            final = os.path.join(work, f"{base}.dwg")
            shutil.move(produced, final)
            return [final], {}
        return convert_dxf(src, out, opts, work, base)
    if real == 'dwg':
        dxf = _dwg_to_dxf(src, work)
        return convert_dxf(dxf, out, opts, work, base)
    if real in DRAW_IN:
        return convert_draw(src, in_ext, out, opts, work, base)
    if real in ('csv',) or group == 'spreadsheet':
        return convert_spreadsheet(src, in_ext, out, opts, work, base)
    if group == 'data':
        return convert_data(src, in_ext, out, opts, work, base)
    if group == 'document':
        return convert_document(src, in_ext, out, opts, work, base)
    if group == 'presentation':
        return convert_slides(src, in_ext, out, opts, work, base)
    if group in ('audio', 'video'):
        return convert_media(src, in_ext, out, opts, work, base)
    if group == 'archive':
        return convert_archive(src, in_ext, out, opts, work, base)
    if group == 'font':
        return convert_font(src, in_ext, out, opts, work, base)
    if group == 'model3d':
        return convert_model(src, in_ext, out, opts, work, base)
    raise ConversionError(f"Formato de entrada .{in_ext} não suportado.")


def convert_file(src_path, original_name, out, opts, dest_dir=None):
    """Converte um arquivo já salvo em disco. Devolve dict com os nomes gerados
    (na pasta Downloads), avisos e, para saídas vetoriais, uma prévia em SVG."""
    out = (out or '').lower().strip().lstrip('.')
    out = ALIASES.get(out, out)
    base, claimed = split_ext(original_name)
    base = _safe_base(base)
    in_ext = sniff_ext(src_path, claimed)
    if not in_ext:
        raise ConversionError("Não foi possível identificar o formato do arquivo.")

    routes = _build_routes()
    if out not in routes.get(in_ext, {}):
        note = unavailable_notes().get(out) or unavailable_notes().get(in_ext)
        if note:
            raise ConversionError(note)
        raise ConversionError(f"Conversão de {in_ext.upper()} para {out.upper()} não é suportada.")

    work = tempfile.mkdtemp(prefix='conversor_')
    try:
        outputs, extras = _dispatch(src_path, in_ext, out, dict(opts or {}), work, base)
        outputs = [p for p in outputs if p and os.path.exists(p)]
        if not outputs:
            raise ConversionError("A conversão não gerou nenhum arquivo.")

        dest_dir = dest_dir or _downloads_dir()
        if len(outputs) > 1:
            # Várias saídas (uma por página, por aba...) vão num ZIP só
            bundle = os.path.join(work, f"{base}_{out}.zip")
            with zipfile.ZipFile(bundle, 'w', zipfile.ZIP_DEFLATED) as zf:
                for p in outputs:
                    zf.write(p, os.path.basename(p))
            final_src = bundle
            count = len(outputs)
        else:
            final_src = outputs[0]
            count = 1
        target = _unique_path(dest_dir, os.path.basename(final_src))
        shutil.move(final_src, target)
        preview = extras.get('preview_svg')
        if preview and len(preview) > 3_000_000:
            preview = None
        return {
            'filename': os.path.basename(target),
            'path': target,
            'size': os.path.getsize(target),
            'count': count,
            'input_format': in_ext,
            'output_format': out,
            'notes': extras.get('notes', []),
            'preview_svg': preview,
        }
    finally:
        shutil.rmtree(work, ignore_errors=True)


# ---------------------------------------------------------------------------
# Integração com as rotas Flask
# ---------------------------------------------------------------------------

# Só arquivos gerados nesta execução podem ser baixados pela rota de download:
# o token aponta para o caminho real, então nunca se lê um caminho vindo da URL.
_OUTPUTS = {}
_OUTPUTS_LOCK = threading.Lock()


def _register_output(path):
    token = secrets.token_urlsafe(16)
    with _OUTPUTS_LOCK:
        _OUTPUTS[token] = path
    return token


def output_path(token):
    with _OUTPUTS_LOCK:
        path = _OUTPUTS.get(token or '')
    return path if path and os.path.isfile(path) else None


def _save_upload(file_storage, work):
    _, ext = split_ext(file_storage.filename or '')
    path = os.path.join(work, 'entrada' + ('.' + ext if ext else ''))
    file_storage.save(path)
    if os.path.getsize(path) == 0:
        raise ConversionError("O arquivo enviado está vazio.")
    return path


def convert_upload(file_storage, out, opts):
    if not file_storage or not file_storage.filename:
        raise ConversionError("Selecione um arquivo.")
    work = tempfile.mkdtemp(prefix='conversor_up_')
    try:
        src = _save_upload(file_storage, work)
        result = convert_file(src, file_storage.filename, out, opts)
    finally:
        shutil.rmtree(work, ignore_errors=True)
    result['token'] = _register_output(result.pop('path'))
    return result


def detect_upload(file_storage):
    """Descobre o formato real (arquivo sem extensão, .enc do WhatsApp, .dat...)."""
    work = tempfile.mkdtemp(prefix='conversor_det_')
    try:
        src = _save_upload(file_storage, work)
        _, claimed = split_ext(file_storage.filename or '')
        return sniff_ext(src, claimed)
    finally:
        shutil.rmtree(work, ignore_errors=True)


def trace_preview(file_storage, opts):
    """Prévia rápida da vetorização (resolução menor, nada é salvo) para o
    usuário ajustar limiar/modo/cores antes de gerar o DXF/SVG de verdade."""
    if not file_storage or not file_storage.filename:
        raise ConversionError("Selecione um arquivo.")
    work = tempfile.mkdtemp(prefix='conversor_prev_')
    try:
        src = _save_upload(file_storage, work)
        opts = dict(opts or {})
        opts['_stacked'] = opts.get('target') != 'dxf'
        opts['_max_side'] = 1000
        tr = trace_image(src, opts)
        return {'svg': trace_to_svg(tr), 'count': tr.count(), 'layers': len(tr.layers)}
    finally:
        shutil.rmtree(work, ignore_errors=True)
